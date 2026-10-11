"""Offline synthetic timing only; never called by production extraction."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import statistics
import subprocess
import sys
import tempfile
import time


def worker(skill, cases, output, repeats):
    sys.path.insert(0, str(Path(skill) / "scripts"))
    import uwc
    import quality
    import pymupdf
    from openpyxl import load_workbook
    results = []
    cases = Path(cases)
    output = Path(output)
    is_new = hasattr(uwc, "inspect_pages")
    for case_path in sorted(cases.glob("*.json")):
        source = json.loads(case_path.read_text(encoding="utf-8"))
        dataset = source["dataset"]
        def supported_fields(value, schema):
            if '$ref' in schema:
                schema=quality.SCHEMA['$defs'][schema['$ref'].rsplit('/',1)[-1]]
            if isinstance(value,dict) and 'properties' in schema:
                return {k:supported_fields(v,schema['properties'][k]) for k,v in value.items() if k in schema['properties']}
            if isinstance(value,list) and 'items' in schema:return [supported_fields(v,schema['items']) for v in value]
            return value
        dataset=supported_fields(dataset,quality.SCHEMA)
        supported = quality.SCHEMA["properties"]["schema_version"].get("enum", ["1.0.0"])
        dataset["schema_version"] = dataset["vocabulary_version"] = supported[-1]
        if not is_new:
            dataset["schema_version"] = dataset["vocabulary_version"] = "1.0.0"
            for e in dataset["evidence"]:
                e.pop("source_role", None)
                e.pop("expected_series", None)
        durations, prep_times, export_times = [], [], []
        for iteration in range(repeats + 1):
            # Iteration zero warms the same imports and library paths for each case.
            root = output / case_path.stem / str(iteration)
            start = time.perf_counter()
            package_dir = uwc.prepare(cases / source["pdf"], dataset["paper"]["paper_id"], root)
            if is_new:
                uwc.inspect_pages(package_dir, source["selected_pages"], vectors=True)
            prepared = time.perf_counter()
            result = uwc.export(dataset, package_dir)
            finished = time.perf_counter()
            if iteration:
                durations.append(finished - start)
                prep_times.append(prepared - start)
                export_times.append(finished - prepared)
        wb = load_workbook(package_dir / f"{dataset['paper']['paper_id']}_UWC.xlsx", read_only=True)
        open_reviews = [r for r in result["review_queue"] if r["status"] == "open"]
        row = {"case": case_path.stem, "corpus_type": "synthetic", "curves": len(result["curves"]),
               "points": len(result["uwc_points"]), "runtime_median_seconds": statistics.median(durations),
               "preparation_median_seconds": statistics.median(prep_times), "qa_export_median_seconds": statistics.median(export_times),
               "runtime_samples_seconds": durations, "open_reviews": len(open_reviews),
               "critical_reviews": sum(quality.review_severity(r) == "critical" for r in open_reviews) if is_new else None,
               "important_reviews": sum(quality.review_severity(r) == "important" for r in open_reviews) if is_new else None,
               "expected_curve_count": result.get("coverage", {}).get("expected_curve_count"),
               "coverage_percentage": result.get("coverage", {}).get("coverage_percentage"),
               "sheet_count": len(wb.sheetnames),
               "data_column_count": 2 if "UWC_SFCC" in wb.sheetnames else wb["UWC_Data" if is_new else "UWC_Points"].max_column,
               "screenshots": len(list(package_dir.rglob("*.png"))), "selected_pages": source["selected_pages"],
               "primary_source": source["primary_source"], "identified_sources": source["identified_sources"],
               "point_fingerprint": hashlib.sha256(json.dumps(result["uwc_points"], sort_keys=True, ensure_ascii=False).encode()).hexdigest()}
        wb.close()
        results.append(row)
    (output / "metrics.json").write_text(json.dumps(results, indent=2), encoding="utf-8")


def make_cases(out):
    import pymupdf as fitz
    import numpy as np
    from test_extractor import fixture, blank
    out.mkdir(parents=True, exist_ok=True)
    configurations = [("single_curve", 1, ["freezing"], False, False),
                      ("freezing_thawing", 1, ["freezing", "thawing"], False, False),
                      ("multiple_conditions", 5, ["freezing", "thawing"], False, False),
                      ("measured_fitted", 1, ["freezing"], True, False),
                      ("complex_figure", 3, ["freezing"], False, True)]
    for name, condition_count, branches, fitted, complex_figure in configurations:
        data = fixture()
        data["derivations"] = []
        base_curve, base_point, base_condition = data["curves"][0], data["uwc_points"][0], data["conditions"][0]
        data["curves"], data["uwc_points"], data["conditions"] = [], [], []
        data["paper"].update(paper_title=f"SYNTHETIC TEST ONLY: {name}", pdf_filename=name + ".pdf")
        data["samples"][0].update(soil_type_normalized="other", soil_type_normalization_status="unmapped")
        source = data["evidence"][0]
        source.update(source_type="figure", figure_number="9", page=9, table_number=None, table_caption=None,
                      figure_caption="Synthetic primary measured SFCC curves (software test; not a research paper).", expected_series=[])
        for ci in range(1, condition_count + 1):
            cond = dict(base_condition, condition_id=f"cond_{ci}", condition_label_reported=f"Initial suction {ci * 100} kPa")
            data["conditions"].append(cond)
            for branch in branches:
                cid = f"c_{ci}_{branch}"
                curve = dict(base_curve, curve_id=cid, condition_id=cond["condition_id"], branch=branch,
                             uwc_basis="volumetric", extraction_method="pdf_vector", figure_id="9", page=9,
                             digitization=blank("digitization", coordinate_system="image_pixels", x_axis_scale="linear",
                                 y_axis_scale="linear", x_axis_orientation="normal", y_axis_orientation="normal",
                                 x_calibration_points=4, y_calibration_points=4, x_calibration_residual_px=.1, y_calibration_residual_px=.1,
                                 axis_confidence="low" if complex_figure and ci == 1 else "high",
                                 calibration_artifact=".internal/calibration.json", benchmark_status="pending"))
                data["curves"].append(curve)
                source["expected_series"].append({"condition_id": cond["condition_id"], "branch": branch, "curve_data_type": "measured_point"})
                for pi, temperature in enumerate(np.linspace(-20, 0, 13), 1):
                    value = float(.1 + .3 * np.exp(temperature / 4))
                    data["uwc_points"].append(dict(base_point, point_id=f"pt_{ci}_{branch}_{pi}", curve_id=cid, point_index=pi,
                        temperature_celsius=float(temperature), temperature_original_value=float(temperature), temperature_value_source="figure_digitized",
                        uwc_original_value=value, uwc_original_unit="cubic_meter_per_cubic_meter", uwc_original_basis="volumetric", uwc_original_unit_reported="m³/m³",
                        uwc_preferred_value=value, uwc_preferred_value_source="figure_digitized", uwc_derivation_id=None,
                        value_source="figure_digitized", extraction_method="pdf_vector"))
        validation = dict(source, evidence_id="ev_validation", figure_number="12", figure_panel="a", page=12,
                          source_role="validation", expected_series=[], figure_caption="Synthetic supporting values at three temperatures.")
        data["evidence"].append(validation)
        selected = [9]
        if fitted:
            fitting_source = dict(source, evidence_id="ev_fit", figure_number="10", page=10, source_role="sfcc_model", expected_series=[], figure_caption="Synthetic fitted SFCC.")
            fitting_source["expected_series"] = [{"condition_id": "cond_1", "branch": "freezing", "curve_data_type": "fitted_curve"}]
            data["evidence"].append(fitting_source)
            curve = copy.deepcopy(data["curves"][0])
            curve.update(curve_id="c_fit", figure_id="10", page=10, curve_data_type="fitted_curve", evidence_ids=["ev_fit"])
            curve["digitization"]["normalized_error_tolerance"] = .002
            data["curves"].append(curve)
            for pi, temperature in enumerate(np.linspace(-20, 0, 100), 1):
                value = float(.1 + .3 * np.exp(temperature / 4))
                data["uwc_points"].append(dict(data["uwc_points"][0], point_id=f"fit_{pi}", curve_id="c_fit", point_index=pi,
                    curve_data_type="fitted_curve", evidence_ids=["ev_fit"], temperature_celsius=float(temperature),
                    temperature_original_value=float(temperature), uwc_original_value=value, uwc_preferred_value=value))
            selected.append(10)
        # Non-critical terminology that the previous validator sent to review.
        for index in range(4):
            prop = dict(data["soil_properties"][0], property_id=f"mineral_{index}", property="mineral_fraction",
                        property_reported="mineral fraction", value=index + 1, original_value=index + 1,
                        unit="percent", original_unit="percent", unit_reported="%", mineral_name_reported=f"Synthetic mineral {index}", mineral_name_normalized="other")
            data["soil_properties"].append(prop)
        pdf = out / f"{name}.pdf"
        with fitz.open() as doc:
            for number in range(1, 13):
                page = doc.new_page(width=500, height=700)
                page.insert_text((40, 40), f"SYNTHETIC REGRESSION TEST {name}; page {number}")
                page.insert_text((40, 60), "UWC freezing thawing SFCC" if number in selected else "Unrelated material for preprocessing cost")
                for index in range(30):
                    page.draw_line((40 + index * 5, 120), (190 + index * 5, 300 + index), color=(.2, .3, .6))
                page.insert_text((40, 400), "Synthetic test values; no real-paper accuracy claim.")
            doc.save(pdf)
        data["paper"]["source_sha256"] = hashlib.sha256(pdf.read_bytes()).hexdigest()
        value = {"pdf": pdf.name, "dataset": data, "selected_pages": selected,
                 "primary_source": "Fig. 9, p. 9", "identified_sources": ["Fig. 9", "Fig. 12a"] + (["Fig. 10"] if fitted else [])}
        (out / f"{name}.json").write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--baseline")
    p.add_argument("--out", required=True)
    p.add_argument("--repeats", type=int, default=3)
    p.add_argument("--worker")
    p.add_argument("--cases")
    args = p.parse_args()
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    if args.worker:
        worker(args.worker, args.cases, out, args.repeats)
        return
    if not args.baseline or not (Path(args.baseline) / "scripts/uwc.py").is_file():
        p.error("--baseline must name a preserved old skill directory")
    make_cases(out / "cases")
    for version, skill in (("old", Path(args.baseline).resolve()), ("new", Path(__file__).resolve().parents[1])):
        subprocess.run([sys.executable, "-X", "utf8", str(Path(__file__).resolve()), "--worker", str(skill),
                        "--cases", str(out / "cases"), "--out", str(out / version), "--repeats", str(args.repeats)], check=True)
    metrics = {v: json.loads((out / v / "metrics.json").read_text()) for v in ("old", "new")}
    for old, new in zip(metrics["old"], metrics["new"]):
        assert old["point_fingerprint"] == new["point_fingerprint"], "Numeric/provenance records changed"
    metrics["scope"] = "Synthetic PDF preparation (including selected vector extraction in new version) plus QA/export; excludes semantic extraction and digitization accuracy. Median of warm runs."
    (out / "comparison.json").write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    print(out / "comparison.json")


if __name__ == "__main__":
    main()
