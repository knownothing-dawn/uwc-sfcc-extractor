"""Paired-curve export checks using synthetic observations, not paper evidence."""
import copy
from collections import Counter
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from test_extractor import fixture, blank
from presentation import curve_blocks, point_pairs, soil_table, workbook
from design import property_cell, test_condition_table


def multi_curve_fixture():
    data = fixture()
    base_c, base_p, base_state = data["curves"][0], data["uwc_points"][0], data["conditions"][0]
    data["curves"], data["uwc_points"], data["conditions"] = [], [], []
    data["soil_properties"], data["derivations"] = [], []
    for index, suction in enumerate((57, 100, 200, 300, 400)):
        state = dict(base_state, condition_id=f"state_{index}", cycle_number=None,
                     condition_label_reported=f"Initial matric suction = {suction} kPa")
        data["conditions"].append(state)
        for branch, count in (("freezing", 12), ("thawing", 11)):
            curve = dict(base_c, curve_id=f"series_{index}_{branch}", condition_id=state["condition_id"],
                         cycle_number=None, branch=branch, uwc_basis="volumetric")
            data["curves"].append(curve)
            for i in range(count):
                temperature = -i if branch == "freezing" else i - count
                value = .1 + index * .02 + i * .001
                data["uwc_points"].append(dict(base_p, point_id=f"obs_{index}_{branch}_{i}", curve_id=curve["curve_id"],
                    point_index=i + 1, temperature_celsius=temperature, uwc_original_value=value,
                    uwc_original_unit="cubic_meter_per_cubic_meter", uwc_original_basis="volumetric",
                    uwc_preferred_value=value, uwc_preferred_value_source="reported", uwc_derivation_id=None))
    data["evidence"][0]["expected_series"] = [{"condition_id":c["condition_id"], "branch":c["branch"],
        "curve_data_type":c["curve_data_type"]} for c in data["curves"]]
    return data


class PresentationTests(unittest.TestCase):
    def test_ten_blocks_115_pairs_bands_four_four_two_and_no_internal_ids(self):
        data = multi_curve_fixture()
        before = copy.deepcopy(data)
        expected = curve_blocks(data)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paired.xlsx"
            workbook(data, path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["UWC_SFCC"])
            ws = wb.active
            headers = [c for row in ws for c in row if c.value == "Temp (°C)"]
            self.assertEqual(list(Counter(c.row for c in headers).values()), [4, 4, 2])
            actual = []
            for header, block in zip(headers, expected):
                self.assertEqual(ws.cell(header.row - 2, header.column).value, block["title"])
                self.assertEqual(ws.cell(header.row - 1, header.column).value, block["subtitle"])
                self.assertEqual(ws.cell(header.row, header.column + 1).value, "θu (m³/m³)")
                pairs = []
                r = header.row + 1
                while isinstance(ws.cell(r, header.column).value, (int, float)):
                    temperature, value = ws.cell(r, header.column).value, ws.cell(r, header.column + 1).value
                    self.assertIsInstance(value, (int, float))
                    pairs.append((temperature, value))
                    r += 1
                self.assertEqual(len(pairs), len(block["pairs"]))
                for actual_pair, expected_pair in zip(pairs, block["pairs"]):
                    for value, source in zip(actual_pair, expected_pair):
                        self.assertAlmostEqual(value, source, places=14)
                actual.extend(pairs)
            self.assertEqual(len(actual), 115)
            visible = {str(c.value) for row in ws for c in row if c.value is not None}
            self.assertFalse({c["curve_id"] for c in data["curves"]} & visible)
            self.assertFalse({p["point_id"] for p in data["uwc_points"]} & visible)
            self.assertFalse(any(c.data_type == "f" for row in ws for c in row))
            self.assertIn("Basic Soil Properties", visible)
            self.assertIn("Test Conditions", visible)
            self.assertEqual(len(soil_table(data)[1]), 1)
            wb.close()
        self.assertEqual(data, before)

    def test_recorded_point_order_overrides_temperature_sort_and_keeps_duplicates(self):
        data = fixture()
        base = data["uwc_points"][0]
        points = [dict(base, point_index=2, temperature_celsius=-1, uwc_preferred_value=.2),
                  dict(base, point_index=1, temperature_celsius=-1, uwc_preferred_value=.1),
                  dict(base, point_index=3, temperature_celsius=-8, uwc_preferred_value=.3)]
        self.assertEqual(point_pairs(points, "freezing")[1], [(-1, .1), (-1, .2), (-8, .3)])

    def test_fallback_order_depends_on_branch(self):
        base = fixture()["uwc_points"][0]
        points = [dict(base, point_index=None, temperature_celsius=t) for t in (-1, -8, -3)]
        self.assertEqual([t for t, _ in point_pairs(points, "freezing")[1]], [-1, -3, -8])
        self.assertEqual([t for t, _ in point_pairs(points, "thawing")[1]], [-8, -3, -1])

    def test_soil_units_scopes_and_conflicts(self):
        data = fixture()
        data["conditions"][0]["condition_label_reported"] = "State 1"
        data["conditions"].append(dict(data["conditions"][0], condition_id="cond_2", condition_label_reported="State 2"))
        density = data["soil_properties"][0]
        density["condition_id"] = None
        data["soil_properties"].append(dict(density, property_id="rho_2", condition_id="cond_2", value=1900))
        for key, value, unit in [("specific_surface_area", 49800, "square_meter_per_kilogram"),
                                 ("porosity", 28.2, "percent"), ("degree_of_saturation", 73, "percent"),
                                 ("initial_volumetric_water_content", 20.6, "percent"),
                                 ("liquid_limit", .3877, "fraction")]:
            data["soil_properties"].append(dict(density, property_id=key, property=key, value=value, unit=unit))
        before = copy.deepcopy(data)
        headers, rows = soil_table(data)
        for header, expected in (("SSA (m²/g)", 49.8), ("Porosity (-)", .282), ("LL (%)", 38.77)):
            self.assertAlmostEqual(rows[0][headers.index(header)], expected)
        self.assertEqual(len(rows),1)
        self.assertNotIn("Test dry density (g/cm³)",headers)
        state_headers,state_rows=test_condition_table(data)
        self.assertEqual(state_rows[0][state_headers.index("Test dry density (g/cm³)")],"1.600, 1.900")
        self.assertEqual(data, before)
        data["soil_properties"].append(dict(density, property_id="conflict", value=1700))
        densities=[p for p in data["soil_properties"] if p["property"]=="dry_density" and p["condition_id"] is None]
        self.assertEqual(property_cell(densities,"dry_density","gram_per_cubic_centimeter"), "Conflict; see report")

    def test_ambiguous_density_never_fills_dry_density(self):
        data = fixture()
        data["soil_properties"][0].update(property="other", property_reported="density", definition_status="uncertain")
        self.assertNotIn("Test dry density (g/cm³)",test_condition_table(data)[0])

    def test_missing_preferred_value_uses_one_original_axis_for_entire_curve(self):
        data = fixture()
        base = data["uwc_points"][0]
        data["uwc_points"].append(dict(base, point_id="point_2", point_index=2, temperature_celsius=-3,
                                      uwc_original_value=10, uwc_original_unit="percent", uwc_preferred_value=None))
        block = curve_blocks(data)[0]
        self.assertEqual(block["uwc_heading"], "w_u, dry mass (kg/kg)")
        self.assertEqual(block["pairs"], [(-2, .2), (-3, .1)])

    def test_incompatible_raw_units_are_not_mislabeled_or_silently_dropped(self):
        base = fixture()["uwc_points"][0]
        points = [dict(base, uwc_preferred_value=None), dict(base, point_index=2, uwc_preferred_value=None,
                                                             uwc_original_unit="other", uwc_original_unit_reported="signal")]
        with self.assertRaises(ValueError):
            point_pairs(points, "freezing")

    def test_review_note_is_scoped_and_normal_confidence_is_hidden(self):
        data = multi_curve_fixture()
        data["review_queue"] = [blank("review", status="open", object_id=data["curves"][1]["curve_id"], issue_type="axis_uncertain")]
        notes = [b["note"] for b in curve_blocks(data)]
        self.assertEqual(notes[1], "Review required: axis uncertain.")
        self.assertEqual(sum(n is not None for n in notes), 1)

    def test_sample_and_data_type_are_separate_and_identifiable(self):
        data = multi_curve_fixture()
        sample = dict(data["samples"][0], sample_id="sample_other", sample_name_reported="Second soil")
        state = dict(data["conditions"][0], sample_id=sample["sample_id"], condition_id="state_other")
        fitted = dict(data["curves"][0], curve_id="fitted", curve_data_type="fitted_curve")
        other = dict(data["curves"][0], curve_id="other", sample_id=sample["sample_id"], condition_id=state["condition_id"])
        data["samples"].append(sample)
        data["conditions"].append(state)
        data["curves"].extend([other, fitted])
        blocks = curve_blocks(data)
        self.assertEqual(blocks[-1]["sample_id"], "sample_other")
        self.assertIn("fitted curve", next(b["subtitle"] for b in blocks if b["curve_id"] == "fitted"))

    def test_uncertain_axis_note_survives_saved_workbook(self):
        data = fixture()
        data["curves"][0]["digitization"] = {"axis_confidence": "low"}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "warning.xlsx"
            workbook(data, path)
            wb = load_workbook(path)
            self.assertIn("Review required: axis uncertain.", [c.value for row in wb.active for c in row])
            wb.close()


if __name__ == "__main__":
    unittest.main()
