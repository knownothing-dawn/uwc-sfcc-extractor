"""Synthetic regression tests. These are not literature benchmark results."""
import copy
import json
import math
from pathlib import Path
import sys
import tempfile
import unittest
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from science import normalize_term, normalize_point, make_input, derive, convert_value, convert_record
from quality import SCHEMA, issues, qa, read_json
from geometry import (adaptive_indices, calibrate_axis, normalized_points, segment_distances,
                      marker_centers, trace_columns, vector_polylines, save_overlay)
from uwc import empty_package, export, prepare, inspect_pages, markdown_report
from benchmark import compare, marker_metrics


def blank(name, **values):
    result = {key: None for key in SCHEMA["$defs"][name]["required"]}
    result.update(values)
    return result


def fixture():
    """Synthetic Chinese source, dry-mass UWC and density, one documented conversion."""
    d = empty_package("paper_test", "synthetic.pdf", "0" * 64)
    d["paper"].update(paper_title="冻土未冻水试验（合成测试，非真实论文）", authors=["张三", "李四"], first_author="张三",
                      publication_year=2001, journal="合成测试期刊", doi="synthetic:test", language="zh",
                      citation_short="张三和李四（2001）", extraction_status="complete", coverage_verified=True,
                      uwc_relevance="high", extractable_uwc_data=True, confidence="high", evidence_ids=["ev_1"])
    d["evidence"] = [blank("evidence", evidence_id="ev_1", source_type="table", source_file="original.pdf", page=1,
                           table_caption="表1 未冻水含量（合成数据）", quoted_evidence="干密度为1600 kg/m³，未冻水质量含水率为20%。",
                           source_context="质量含水率以干土质量为基准。", source_language="zh", transcription_status="verified")]
    d["samples"] = [blank("sample", sample_id="s_1", paper_id="paper_test", sample_name_reported="土样甲",
                          soil_type_reported="兰州黄土", soil_type_normalized="loess", soil_type_normalization_status="verified_mapping",
                          material_category="natural_soil", sample_preparation=["remolded"], sample_preparation_reported="重塑土样",
                          confidence="high", evidence_ids=["ev_1"])]
    d["conditions"] = [blank("condition", condition_id="cond_1", sample_id="s_1", cycle_number=1,
                             condition_label_reported="第一次冻结", thermal_history_reported="降温冻结",
                             confidence="high", evidence_ids=["ev_1"])]
    d["soil_properties"] = [blank("property", property_id="rho_1", sample_id="s_1", condition_id="cond_1",
                                 property="dry_density", property_reported="干密度", value=1600, value_reported="1600",
                                 unit="kilogram_per_cubic_meter", unit_reported="kg/m³", original_value=1600,
                                 original_unit="kilogram_per_cubic_meter", definition_reported="干土质量/土样体积",
                                 definition_status="explicit", value_source="reported", confidence="high", evidence_ids=["ev_1"])]
    d["curves"] = [blank("curve", curve_id="c_1", paper_id="paper_test", sample_id="s_1", condition_id="cond_1",
                         page=1, branch_reported="冻结", branch="freezing", branch_source="reported", cycle_number=1,
                         method_reported="核磁共振", method="nuclear_magnetic_resonance", curve_data_type="measured_point",
                         legend_label="土样甲（第一次冻结）", uwc_basis="gravimetric_dry_mass", uwc_definition_reported="未冻水质量/干土质量",
                         temperature_unit="degree_celsius", extraction_method="table_direct", data_origin="original_data",
                         needs_manual_review=False, confidence="high", evidence_ids=["ev_1"])]
    d["uwc_points"] = [blank("point", point_id="pt_1", curve_id="c_1", point_index=1,
                             temperature_celsius=-2, temperature_original_value=-2, temperature_original_unit="degree_celsius",
                             temperature_original_unit_reported="℃", temperature_value_source="reported",
                             uwc_original_value=.2, uwc_original_unit="kilogram_per_kilogram", uwc_original_unit_reported="kg/kg",
                             uwc_original_basis="gravimetric_dry_mass", uwc_preferred_value=.32,
                             uwc_preferred_unit="cubic_meter_per_cubic_meter", uwc_preferred_basis="volumetric",
                             uwc_preferred_value_source="derived", uwc_derivation_id="d_1", curve_data_type="measured_point",
                             extraction_method="table_direct", value_source="reported", confidence="high", evidence_ids=["ev_1"])]
    inputs = [make_input("gravimetric_unfrozen_water_content", .2, "kilogram_per_kilogram", "pt_1", "uwc_original_value", ["ev_1"]),
              make_input("dry_density", 1600, "kilogram_per_cubic_meter", "rho_1", evidence_ids=["ev_1"])]
    d["derivations"] = [derive("volumetric_unfrozen_water_content", inputs, "d_1", allow_default_water_density=True)]
    d["evidence"][0].update(source_role="primary_uwc", table_number="1", expected_series=[
        {"condition_id": "cond_1", "branch": "freezing", "curve_data_type": "measured_point"}])
    d['curves'][0].update(variable_definition_verified=True, source_points_verified=True)
    d['source_inspection'] = dict(point_inventory_checked=True, anchor_search_completed=True,
        property_pass_completed=True, evidence_ids=['ev_1'])
    return d


class LanguageContractTests(unittest.TestCase):
    def test_valid_chinese_evidence(self):
        data = fixture()
        self.assertEqual(issues(data), [])
        self.assertEqual(qa(data)["review_queue"], [])

    def test_non_english_key_and_enum_rejected(self):
        for edit in (lambda x: x["samples"][0].update({"土壤": "黄土"}),
                     lambda x: x["samples"][0].update(soil_type_normalized="黄土"),
                     lambda x: x["curves"][0].update(branch="Freezing")):
            d = fixture()
            edit(d)
            self.assertTrue(issues(d))

    def test_alias_does_not_modify_source(self):
        self.assertEqual(normalize_term("branch", "Cooling")["normalized"], "cooling")
        original = "  黄土  "
        result = normalize_term("soil_type", original)
        self.assertEqual(result["reported"], original)
        self.assertEqual(result["normalized"], "loess")
        self.assertFalse(normalize_term("soil_type", "unrecognized material")["needs_manual_review"])
        self.assertTrue(normalize_term("soil_type", "unrecognized material", affects_uwc=True)["needs_manual_review"])

    def test_uncertain_display_translation_is_flagged_without_action_queue(self):
        d = fixture()
        d["display_translations"] = [{"translation_id": "tr_1", "source_pointer": "/paper/paper_title",
                                       "language": "en", "text": "Synthetic title", "translation_status": "uncertain"}]
        result = qa(d)
        self.assertEqual(result["review_queue"], [])
        self.assertIn("translation is uncertain", markdown_report(result))
        self.assertEqual(result["paper"]["paper_title"], d["paper"]["paper_title"])

    def test_duplicate_json_keys_and_nan_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.json"
            for raw in ('{"a":1,"a":2}', '{"a":NaN}'):
                path.write_text(raw)
                with self.assertRaises(ValueError):
                    read_json(path)


class ScientificTests(unittest.TestCase):
    def test_all_allowed_relations(self):
        def i(k, v, unit):
            return make_input(k, v, unit, "source_" + k, evidence_ids=["ev_1"])
        density = "kilogram_per_cubic_meter"
        cases = [
            ("porosity", [i("dry_density", 1600, density), i("particle_density", 2650, density)], 1 - 1600 / 2650),
            ("void_ratio_from_porosity", [i("porosity", .4, "fraction")], 2 / 3),
            ("void_ratio_from_densities", [i("dry_density", 1600, density), i("particle_density", 2650, density)], 2650 / 1600 - 1),
            ("degree_of_saturation", [i("volumetric_water_content", .3, "cubic_meter_per_cubic_meter"), i("porosity", .4, "fraction")], .75),
            ("wet_density", [i("dry_density", 1600, density), i("gravimetric_water_content", .2, "kilogram_per_kilogram")], 1920),
            ("particle_density", [i("specific_gravity", 2.65, "dimensionless"), i("water_density", 1000, density)], 2650),
        ]
        for relation, inputs, answer in cases:
            with self.subTest(relation=relation):
                self.assertAlmostEqual(derive(relation, inputs, "derivation")["result"], answer)
        result = fixture()["derivations"][0]
        self.assertAlmostEqual(result["result"], .32)
        self.assertTrue(result["assumption_used"])
        self.assertEqual(result["assumptions"][0]["value"], 1000)

    def test_ambiguous_percent_rejected(self):
        with self.assertRaises(ValueError):
            convert_value(20, "percent", "kilogram_per_kilogram")
        self.assertEqual(convert_value(20, "percent", "kilogram_per_kilogram", "gravimetric_dry_mass"), .2)
        with self.assertRaises(ValueError):
            convert_value(.2, "kilogram_per_kilogram", "cubic_meter_per_cubic_meter", "volumetric")

    def test_unit_conversions(self):
        cases = [(273.15, "kelvin", "degree_celsius", 0), (32, "degree_fahrenheit", "degree_celsius", 0),
                 (1.6, "gram_per_cubic_centimeter", "kilogram_per_cubic_meter", 1600),
                 (2000, "square_meter_per_kilogram", "square_meter_per_gram", 2),
                 (1000, "microsiemens_per_centimeter", "decisiemens_per_meter", 1),
                 (2, "degree_celsius_per_minute", "degree_celsius_per_hour", 120)]
        for value, source, target, answer in cases:
            self.assertAlmostEqual(convert_value(value, source, target), answer)

    def test_uncertain_input_and_invalid_density_rejected(self):
        inputs = fixture()["derivations"][0]["input_variables"][:2]
        inputs[0]["definition_status"] = "uncertain"
        with self.assertRaises(ValueError):
            derive("volumetric_unfrozen_water_content", inputs, "d", allow_default_water_density=True)
        inputs[0]["definition_status"] = "explicit"
        inputs[1]["value"] = 0
        with self.assertRaises(ValueError):
            derive("volumetric_unfrozen_water_content", inputs, "d", allow_default_water_density=True)

    def test_tampered_derivation_detected(self):
        d = fixture()
        d["derivations"][0]["result"] = .5
        d["uwc_points"][0]["uwc_preferred_value"] = .5
        self.assertTrue(any("incorrect calculated result" in s for s in issues(d)))

    def test_wrong_binding_and_provenance_rejected(self):
        d = fixture()
        d["curves"][0]["cycle_number"] = 2
        self.assertTrue(issues(d))
        d = fixture()
        d["uwc_points"][0]["value_source"] = "figure_digitized"
        self.assertTrue(issues(d))

    def test_kelvin_has_separate_derivation(self):
        d = fixture()
        pt = d["uwc_points"][0]
        pt.update(temperature_original_value=271.15, temperature_original_unit="kelvin", temperature_original_unit_reported="K",
                  temperature_value_source="derived", temperature_derivation_id="d_t")
        item = make_input("temperature_celsius", 271.15, "kelvin", "pt_1", "temperature_original_value", ["ev_1"])
        d["derivations"].append(convert_record(item, "degree_celsius", "d_t"))
        self.assertEqual(issues(d), [])

    def test_original_percent_preserved_during_mass_to_volume_conversion(self):
        d = fixture()
        pt = d["uwc_points"][0]
        pt.update(uwc_original_value=20, uwc_original_unit="percent", uwc_original_unit_reported="%")
        inputs = d["derivations"][0]["input_variables"][:2]
        inputs[0].update(value=20, unit="percent")
        d["derivations"] = [derive("volumetric_unfrozen_water_content", inputs, "d_1", allow_default_water_density=True)]
        self.assertEqual(issues(d), [])
        self.assertAlmostEqual(d["derivations"][0]["result"], .32)

    def test_derivation_cannot_attach_to_different_condition(self):
        d = fixture()
        condition = copy.deepcopy(d["conditions"][0])
        condition["condition_id"] = "cond_2"
        d["conditions"].append(condition)
        curve = copy.deepcopy(d["curves"][0])
        curve.update(curve_id="c_2", condition_id="cond_2")
        d["curves"].append(curve)
        point = copy.deepcopy(d["uwc_points"][0])
        point.update(point_id="pt_2", curve_id="c_2")
        d["uwc_points"].append(point)
        self.assertTrue(any("mixes sample or condition" in e for e in issues(d)))

    def test_conflicts_are_retained(self):
        d = fixture()
        d["uwc_points"][0].update(uwc_preferred_value=None, uwc_preferred_unit=None, uwc_preferred_basis=None,
                                 uwc_preferred_value_source=None, uwc_derivation_id=None)
        d["derivations"] = []
        duplicate = copy.deepcopy(d["soil_properties"][0])
        duplicate.update(property_id="rho_2", value=1480, original_value=1480, value_reported="1480")
        d["soil_properties"].append(duplicate)
        result = qa(d)
        self.assertEqual([r["value"] for r in result["soil_properties"]], [1600, 1480])
        self.assertIn("conflicting_values", [r["issue_type"] for r in result["review_queue"]])

    def test_qa_does_not_promote_human_approval(self):
        d = fixture();d['paper']['review_status'] = 'pending'
        result = qa(d)
        self.assertEqual(result["paper"]["legacy_review_status"], "pending")
        self.assertEqual(result["paper"]["review_status"], "database_ready")


class GeometryTests(unittest.TestCase):
    def test_linear_reversed_and_log_calibration(self):
        axis = calibrate_axis([[0, 0], [50, -10], [100, -20]], "linear", .1)
        self.assertAlmostEqual(float(axis.to_data(25)), -5)
        self.assertAlmostEqual(float(axis.to_coordinate(-5)), 25)
        log = calibrate_axis([[0, 1], [50, 10], [100, 100]], "log", .1)
        self.assertAlmostEqual(float(log.to_data(25)), math.sqrt(10))

    def test_bad_ticks_and_special_axes_rejected(self):
        for anchors, scale in (([[0, 0], [100, 1]], "linear"), ([[0, 0], [50, 7], [100, 10]], "linear"),
                               ([[0, -10], [50, -1], [100, 0]], "log"), ([[0, 0], [50, 1], [100, 2]], "broken")):
            with self.assertRaises(ValueError):
                calibrate_axis(anchors, scale, .1)

    def test_adaptive_preserves_endpoints_turns_and_error_bound(self):
        x = np.linspace(-20, 0, 1001)
        y = .4 * np.exp(x / 2)
        points = np.column_stack((x, y))
        indices = adaptive_indices(points, [-20, 0, 0, .4], .002)
        self.assertEqual((indices[0], indices[-1]), (0, 1000))
        self.assertLess(len(indices), 100)
        normalized = normalized_points(points, [-20, 0, 0, .4])
        for start, end in zip(indices[:-1], indices[1:]):
            self.assertLessEqual(max(segment_distances(normalized[start:end + 1], normalized[start], normalized[end])), .002 + 1e-12)
        turning = [[0, 0], [1, 1], [2, 1], [3, 0]]
        self.assertEqual(adaptive_indices(turning, [0, 3, 0, 1], 10), [0, 1, 2, 3])

    def test_isolated_markers_are_not_resampled(self):
        mask = np.zeros((50, 100), bool)
        mask[10:13, 20:23] = True
        mask[30:33, 70:73] = True
        result = marker_centers(mask, 5, 15)
        self.assertEqual(result["centers"], [[21., 11.], [71., 31.]])

    def test_trace_refuses_gaps_and_multiple_runs(self):
        mask = np.zeros((20, 20), bool)
        mask[5:7, 2:12] = True
        self.assertEqual(len(trace_columns(mask)), 10)
        mask[:, 7] = False
        with self.assertRaises(ValueError):
            trace_columns(mask)
        mask[5:7, 7] = True
        mask[12, 7] = True
        with self.assertRaises(ValueError):
            trace_columns(mask)

    def test_vector_controls_are_flattened_and_gaps_preserved(self):
        paths = vector_polylines([("c", [0, 0], [0, 1], [1, 1], [1, 0]), ("l", [2, 0], [3, 1])], .01)
        self.assertEqual(len(paths), 2)
        self.assertGreater(len(paths[0]), 4)
        self.assertEqual(paths[0][0], [0., 0.])
        self.assertEqual(paths[0][-1], [1., 0.])

    def test_benchmark_reports_three_strategies_without_validating_production(self):
        x = np.linspace(-10, 0, 200)
        trace = np.column_stack((x, .4 * np.exp(x / 2)))
        result = compare(trace, trace, [-10, 0, 0, .4], 30, .5, .002)
        self.assertEqual(len(result["results"]), 3)
        self.assertFalse(result["production_threshold_validated"])
        scores = marker_metrics([[0, 0], [1, 1], [.5, .5]], [[0, 0], [1, 1]], [0, 1, 0, 1], .01)
        self.assertEqual(scores["true_positive"], 2)
        self.assertAlmostEqual(scores["precision"], 2 / 3)


class IntegrationTests(unittest.TestCase):
    def test_json_excel_roundtrip_preserves_original_strings(self):
        from openpyxl import load_workbook
        d = fixture()
        d["evidence"][0]["quoted_evidence"] = "=θu（原式） + 未冻水"
        d["paper"]["paper_title"] = "=合成论文标题"
        with tempfile.TemporaryDirectory() as tmp:
            result = export(d, tmp)
            saved = read_json(Path(tmp) / ".internal/paper_test_UWC.json")
            self.assertEqual(saved["evidence"][0]["quoted_evidence"], d["evidence"][0]["quoted_evidence"])
            self.assertEqual(saved["paper"]["paper_title"], d["paper"]["paper_title"])
            wb = load_workbook(Path(tmp) / "paper_test_UWC.xlsx")
            self.assertEqual(wb.sheetnames, ["UWC_SFCC"])
            ws = wb["UWC_SFCC"]
            cell = ws["B2"]
            self.assertEqual(cell.value, "=合成论文标题")
            self.assertEqual(cell.data_type, "s")
            self.assertIn("Basic Soil Properties", [c.value for row in ws for c in row])
            self.assertEqual(sum(c.value == "Temp (°C)" for row in ws for c in row), 1)
            self.assertEqual(result["uwc_points"], d["uwc_points"])
            wb.close()

    def test_prepare_pdf_candidates_and_export_empty_package(self):
        import pymupdf as fitz
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "synthetic.pdf"
            with fitz.open() as doc:
                page = doc.new_page(width=300, height=300)
                page.insert_text((20, 30), "SYNTHETIC TEST: unfrozen water content")
                page.draw_line((30, 100), (200, 200), color=(1, 0, 0))
                page.set_rotation(90)
                doc.save(source)
            output = prepare(source, "paper_one", root / "output")
            data = read_json(output / ".internal/paper_one_UWC.json")
            self.assertEqual(issues(data), [])
            candidates = read_json(output / ".internal/candidates.json")
            self.assertTrue(candidates["pages"][0]["uwc_keyword_hit"])
            self.assertFalse(list(output.rglob("*.png")))
            inspect_pages(output, [1], vectors=True)
            self.assertFalse(list(output.rglob("*.png")))
            self.assertEqual(read_json(output / ".internal/pages/page_0001_coordinates.json")["original_rotation_degrees"], 90)
            self.assertGreater(len(read_json(output / ".internal/pages/page_0001_vectors.json")), 0)
            result = export(data, output)
            self.assertNotEqual(result["paper"]["extraction_status"], "complete")
            self.assertEqual({p.name for p in output.iterdir() if p.is_file()}, {"paper_one_UWC.xlsx", "paper_one_report.md"})
            with self.assertRaises(ValueError):
                prepare(source, "paper_one", root / "output")
            with self.assertRaises(ValueError):
                prepare(source, "../escape", root / "output")

    def test_overlay_output(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as tmp:
            a, b = Path(tmp) / "source.png", Path(tmp) / "overlay.png"
            Image.new("RGB", (100, 100), "white").save(a)
            save_overlay(a, b, [[10, 10], [90, 90]], [[10, 10], [90, 90]])
            with Image.open(b) as image:
                self.assertNotEqual(image.getpixel((50, 50)), (255, 255, 255))

    def test_xml_incompatible_source_is_preserved_in_json(self):
        from openpyxl import load_workbook
        d = fixture()
        original = "原文\x00含控制字符"
        d["paper"]["paper_title"] = original
        with tempfile.TemporaryDirectory() as tmp:
            export(d, tmp)
            self.assertEqual(read_json(Path(tmp) / ".internal/paper_test_UWC.json")["paper"]["paper_title"], original)
            wb = load_workbook(Path(tmp) / "paper_test_UWC.xlsx")
            ws = wb["UWC_SFCC"]
            cell = ws["B2"]
            self.assertIn(".internal/paper_test_UWC.json", cell.value)
            wb.close()


if __name__ == "__main__":
    unittest.main()
