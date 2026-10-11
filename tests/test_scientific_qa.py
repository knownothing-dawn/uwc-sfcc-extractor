"""Coverage and variable regressions; real Luo observations are never completed with invented data."""
import copy
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from test_extractor import fixture, blank
from coverage import curve_coverage
from conversion import convert_preferred
from quality import issues, qa
from presentation import point_pairs, workbook
from design import basic_soil_table, test_condition_table
from science import derive, make_input


def luo():
    return json.loads((Path(__file__).parent/'fixtures/luo_2024.json').read_text(encoding='utf-8'))


def saturation_fixture(value=.67,unit='fraction'):
    d=fixture();d['derivations']=[]
    curve=d['curves'][0]
    curve.update(uwc_basis='unfrozen_water_saturation',raw_variable='unfrozen_water_saturation',raw_variable_reported='Su',
                 variable_definition_verified=True,conversion_property_ids=['porosity'])
    point=d['uwc_points'][0]
    point.update(uwc_original_basis='unfrozen_water_saturation',uwc_original_value=value,uwc_original_unit=unit,
                 uwc_preferred_value=None,uwc_preferred_unit=None,uwc_preferred_basis=None,uwc_preferred_value_source=None,uwc_derivation_id=None)
    por=copy.deepcopy(d['soil_properties'][0]);por.update(property_id='porosity',property='porosity',property_reported='porosity',
        value=.4,original_value=.4,unit='fraction',original_unit='fraction')
    d['soil_properties']=[por]
    return d


class CoverageTests(unittest.TestCase):
    def test_distinct_cycles_are_inventoried_and_visible(self):
        d=fixture();d['conditions'][0]['cycle_number']=1;d['curves'][0]['cycle_number']=1
        d['conditions'].append(dict(d['conditions'][0],condition_id='cycle2',cycle_number=2))
        d['curves'].append(dict(d['curves'][0],curve_id='curve2',condition_id='cycle2',cycle_number=2))
        d['uwc_points'].append(dict(d['uwc_points'][0],point_id='point2',curve_id='curve2',
            uwc_preferred_value=None,uwc_preferred_unit=None,uwc_preferred_basis=None,
            uwc_preferred_value_source=None,uwc_derivation_id=None))
        d['experimental_cases']=[dict(case_id='cycles',sample_id=d['samples'][0]['sample_id'],
            case_label_reported='Freeze–thaw cycles',varying_property='cycle_number',factor_label_reported='Freeze–thaw cycle',
            variables=[dict(property='cycle_number',values=[1,2],unit='dimensionless',values_reported=['1','2'],evidence_ids=['ev_1'])],
            condition_ids=['cond_1','cycle2'],branches=['freezing'],curve_data_types=['measured_point'],
            temperature_range_celsius=None,evidence_ids=['ev_1'])]
        result=qa(d)
        self.assertEqual(result['coverage']['expected_curve_count'],2)
        self.assertEqual(result['coverage']['extracted_curve_count'],2)
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'cycles.xlsx';workbook(result,path);wb=load_workbook(path)
            values=[c.value for row in wb.active for c in row]
            self.assertIn('Cycle = 1',values);self.assertIn('Cycle = 2',values);wb.close()

    @unittest.skipUnless(
        (Path(__file__).parent/'fixtures'/'luo_2024.json').exists(),
        'Private recorded-paper fixture is not distributed in the public repository.',
    )
    def test_real_luo_missing_1240_is_a_failed_completeness_benchmark(self):
        d=luo();d['paper']['extraction_status']='complete'
        result=qa(d);coverage=result['coverage']
        self.assertEqual(coverage['expected_curve_count'],19)
        self.assertEqual(coverage['extracted_curve_count'],18)
        self.assertEqual(coverage['source_occurrence_count'],21)
        self.assertAlmostEqual(coverage['coverage_percentage'],100*18/19)
        self.assertEqual([m['condition_id'] for m in coverage['missing_conditions']],['cond_silt_0p870_1p240'])
        self.assertEqual(result['paper']['extraction_status'],'partial')
        self.assertEqual(len(result['uwc_points']),149)
        self.assertTrue(any(r['severity']=='critical' and '1.240' in r['description'] for r in result['review_queue']))

    @unittest.skipUnless(
        (Path(__file__).parent/'fixtures'/'luo_2024.json').exists(),
        'Private recorded-paper fixture is not distributed in the public repository.',
    )
    def test_dropping_expected_condition_cannot_shrink_the_denominator_silently(self):
        d=luo();d['experimental_cases'][-1]['condition_ids'].remove('cond_silt_0p870_1p240')
        self.assertTrue(any('inventory omits' in e for e in issues(d)))

    @unittest.skipUnless(
        (Path(__file__).parent/'fixtures'/'luo_2024.json').exists(),
        'Private recorded-paper fixture is not distributed in the public repository.',
    )
    def test_source_conflict_requires_explicit_design_alias(self):
        d=luo();next(c for c in d['conditions'] if c['condition_id']=='cond_silt_0p652_0p180').pop('design_aliases')
        self.assertTrue(any('design values' in e or 'inventory omits' in e for e in issues(d)))

    @unittest.skipUnless(
        (Path(__file__).parent/'fixtures'/'luo_2024.json').exists(),
        'Private recorded-paper fixture is not distributed in the public repository.',
    )
    def test_resolving_review_does_not_override_missing_curve_gate(self):
        d=qa(luo())
        for r in d['review_queue']:
            if r['severity']=='critical':r.update(status='resolved',resolution_notes='Acknowledged as absent in public sources')
        d['paper']['extraction_status']='complete'
        self.assertNotEqual(qa(d)['paper']['extraction_status'],'complete')

    def test_identical_duplicate_occurrences_are_visible_once(self):
        d=fixture()
        curve=dict(d['curves'][0],curve_id='duplicate')
        point=dict(d['uwc_points'][0],point_id='duplicate_point',curve_id='duplicate',uwc_derivation_id=None,
                   uwc_preferred_value=None,uwc_preferred_unit=None,uwc_preferred_basis=None,uwc_preferred_value_source=None)
        d['curves'].append(curve);d['uwc_points'].append(point)
        result=qa(d)
        self.assertEqual(result['coverage']['extracted_curve_count'],1)
        self.assertTrue(result['coverage']['duplicate_conditions'][0]['identical'])
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'one.xlsx';workbook(result,path)
            wb=load_workbook(path)
            self.assertEqual(sum(c.value=='Temp (°C)' for row in wb.active for c in row),1);wb.close()

    def test_different_values_are_not_auto_deduplicated(self):
        d=fixture();d['curves'].append(dict(d['curves'][0],curve_id='different'))
        d['uwc_points'].append(dict(d['uwc_points'][0],point_id='different_point',curve_id='different',uwc_original_value=.3,
            uwc_preferred_value=None,uwc_preferred_unit=None,uwc_preferred_basis=None,uwc_preferred_value_source=None,uwc_derivation_id=None))
        result=qa(d)
        self.assertEqual(result['coverage']['extracted_curve_count'],2)
        self.assertFalse(result['coverage']['duplicate_conditions'][0]['identical'])
        self.assertEqual(result['paper']['extraction_status'],'manual_review_required')

    def test_unexpected_condition_is_not_hidden(self):
        d=fixture();d['conditions'].append(dict(d['conditions'][0],condition_id='extra',condition_label_reported='Extra specimen'))
        d['curves'].append(dict(d['curves'][0],curve_id='extra_curve',condition_id='extra'))
        d['uwc_points'].append(dict(d['uwc_points'][0],point_id='extra_point',curve_id='extra_curve',
            uwc_preferred_value=None,uwc_preferred_unit=None,uwc_preferred_basis=None,uwc_preferred_value_source=None,uwc_derivation_id=None))
        result=qa(d)
        self.assertEqual(len(result['coverage']['unexpected_conditions']),1)
        self.assertEqual(result['paper']['extraction_status'],'manual_review_required')


class VariableTests(unittest.TestCase):
    def test_saturation_derivation_preserves_raw_and_uses_matched_porosity(self):
        d=saturation_fixture();before=copy.deepcopy(d);result=convert_preferred(d)
        self.assertEqual(d,before)
        self.assertEqual(issues(result),[])
        point=result['uwc_points'][0]
        self.assertAlmostEqual(point['uwc_preferred_value'],.268)
        self.assertEqual(point['uwc_original_value'],.67)
        self.assertEqual(point['uwc_preferred_value_source'],'derived')
        self.assertEqual(result['derivations'][0]['formula'],'theta_u = porosity * unfrozen_water_saturation')
        self.assertEqual(convert_preferred(result),result)

    def test_percent_saturation_and_kelvin_do_not_change_semantics(self):
        d=saturation_fixture(67,'percent')
        d['uwc_points'][0].update(temperature_original_value=271.15,temperature_original_unit='kelvin')
        result=convert_preferred(d)
        self.assertAlmostEqual(result['uwc_points'][0]['uwc_preferred_value'],.268)
        self.assertEqual(result['uwc_points'][0]['temperature_celsius'],-2)
        self.assertEqual(len(result['derivations']),1)
        self.assertEqual(issues(result),[])

    def test_without_porosity_excel_says_su(self):
        d=saturation_fixture(67,'percent');d['soil_properties']=[];d['curves'][0]['conversion_property_ids']=[]
        result=convert_preferred(d)
        header,pairs=point_pairs(result['uwc_points'],'freezing')
        self.assertEqual(header,'Su (-)');self.assertEqual(pairs,[(-2,.67)])
        self.assertEqual(result['derivations'],[])

    def test_su_cannot_be_relabeled_volumetric(self):
        d=saturation_fixture();d['curves'][0]['uwc_basis']='volumetric'
        self.assertTrue(any('variable and UWC basis' in e for e in issues(d)))

    def test_unknown_saturation_definition_and_wrong_specimen_block_conversion(self):
        d=saturation_fixture();d['curves'][0]['variable_definition_verified']=False
        self.assertIsNone(convert_preferred(d)['uwc_points'][0]['uwc_preferred_value'])
        d=saturation_fixture();d['soil_properties'][0]['condition_id']='different_state'
        self.assertIsNone(convert_preferred(d)['uwc_points'][0]['uwc_preferred_value'])

    def test_maximum_density_never_substitutes_for_test_density(self):
        d=fixture();p=d['uwc_points'][0]
        p.update(uwc_preferred_value=None,uwc_preferred_unit=None,uwc_preferred_basis=None,uwc_preferred_value_source=None,uwc_derivation_id=None)
        d['derivations']=[];d['soil_properties'][0].update(property='maximum_dry_density',property_reported='Maximum dry density')
        self.assertIsNone(convert_preferred(d)['uwc_points'][0]['uwc_preferred_value'])
        d['soil_properties'][0]['property']='dry_density'
        self.assertTrue(any('maximum dry density' in e for e in issues(d)))

    def test_dry_mass_conversion_records_water_density_assumption(self):
        d=fixture();p=d['uwc_points'][0]
        p.update(uwc_preferred_value=None,uwc_preferred_unit=None,uwc_preferred_basis=None,uwc_preferred_value_source=None,uwc_derivation_id=None)
        d['derivations']=[];result=convert_preferred(d)
        self.assertAlmostEqual(result['uwc_points'][0]['uwc_preferred_value'],.32)
        self.assertEqual(result['derivations'][0]['assumptions'][0]['value'],1000)
        self.assertEqual(issues(result),[])

    def test_small_saturation_overshoot_is_retained_without_fatal_review(self):
        result=qa(convert_preferred(saturation_fixture(1.002)))
        self.assertEqual(result['uwc_points'][0]['uwc_original_value'],1.002)
        self.assertAlmostEqual(result['uwc_points'][0]['uwc_preferred_value'],.4008)
        self.assertFalse(any(r['severity']=='critical' for r in result['review_queue']))

    def test_large_saturation_overshoot_requires_review(self):
        result=qa(convert_preferred(saturation_fixture(1.2)))
        self.assertEqual(result['paper']['extraction_status'],'manual_review_required')

    def test_acknowledging_invalid_saturation_does_not_make_data_complete(self):
        result=qa(convert_preferred(saturation_fixture(1.2)))
        for r in result['review_queue']:
            r.update(status='resolved',resolution_notes='Acknowledged')
        self.assertEqual(qa(result)['paper']['extraction_status'],'manual_review_required')

    def test_unknown_saturation_unit_stays_reviewable(self):
        from reporting import markdown_report
        d=saturation_fixture(unit='unknown')
        result=qa(convert_preferred(d))
        self.assertEqual(result['paper']['extraction_status'],'manual_review_required')
        self.assertIn('Saturation unit cannot',markdown_report(result))
        self.assertIn('review unit',point_pairs(result['uwc_points'],'freezing')[0])

    def test_derived_volumetric_value_above_porosity_is_critical(self):
        d=fixture()
        por=copy.deepcopy(d['soil_properties'][0])
        por.update(property_id='porosity',property='porosity',property_reported='porosity',
                   value=.2,original_value=.2,unit='fraction',original_unit='fraction')
        d['soil_properties'].append(por)
        result=qa(d)
        self.assertTrue(any(r['severity']=='critical' and 'above porosity' in r['description'] for r in result['review_queue']))

    @unittest.skipUnless(
        (Path(__file__).parent/'fixtures'/'luo_2024.json').exists(),
        'Private recorded-paper fixture is not distributed in the public repository.',
    )
    def test_luo_structure_materials_states_and_unique_blocks(self):
        d=luo();self.assertEqual(issues(d),[])
        self.assertEqual(len(basic_soil_table(d)[1]),3)
        self.assertEqual(len(test_condition_table(d)[1]),6)
        self.assertIn('Maximum dry density (g/cm³)',basic_soil_table(d)[0])
        self.assertIn('Test dry density (g/cm³)',test_condition_table(d)[0])
        with tempfile.TemporaryDirectory() as tmp:
            file=Path(tmp)/'luo.xlsx';workbook(d,file);wb=load_workbook(file);ws=wb.active
            cells=[c.value for row in ws for c in row]
            self.assertEqual(cells.count('Temp (°C)'),18)
            self.assertEqual(cells.count('θu (m³/m³)'),18)
            self.assertTrue(any(isinstance(v,str) and 'Missing: C0 = 1.240' in v for v in cells))
            self.assertFalse(any(isinstance(v,str) and 'Cycle' in v for v in cells))
            for title in ['Paper Information','Basic Soil Properties','Test Conditions','UWC/SFCC Data']:
                self.assertIn(title,cells)
            wb.close()


if __name__=='__main__':unittest.main()
