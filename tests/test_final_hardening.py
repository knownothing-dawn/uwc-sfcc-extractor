"""Final hardening regressions; source papers are represented by recorded small fixtures."""
import copy
import json
import math
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from test_extractor import blank
from test_hardening import graphical, run
from design import test_condition_table
from quality import issues
from trace_processing import process_continuous_traces
from uwc import export


ROOT = Path(__file__).parent


def dense_trace(unique=201):
    data = graphical()
    curve = data['curves'][0]
    curve.update(curve_data_type='unknown', data_semantics='measured', representation='continuous_measured_trace',
                 trace_geometry='single_valued', extraction_method='raster_curve')
    data['evidence'][0]['expected_series'][0].update(curve_data_type='unknown', data_semantics='measured',
                                                     representation='continuous_measured_trace')
    data['soil_properties'] = []
    data['derivations'] = []
    points=[]
    for i in range(unique):
        x=-20+20*i/(unique-1)
        base=.10+.20/(1+math.exp(-3.2*(x+1.5)))
        for side,offset in enumerate((-.00025,.00025)):
            points.append(blank('point', point_id=f'raw_{i}_{side}', curve_id='c_1', point_index=len(points)+1,
                trace_order=len(points)+1, temperature_celsius=x, temperature_original_value=x,
                temperature_original_unit='degree_celsius', temperature_original_unit_reported='°C',
                temperature_value_source='figure_digitized', uwc_original_value=base+offset,
                uwc_original_unit='cubic_meter_per_cubic_meter', uwc_original_unit_reported='m³/m³',
                uwc_original_basis='volumetric', uwc_preferred_value=base+offset,
                uwc_preferred_unit='cubic_meter_per_cubic_meter', uwc_preferred_basis='volumetric',
                uwc_preferred_value_source='figure_digitized', curve_data_type='unknown',
                extraction_method='raster_curve', value_source='figure_digitized', confidence='high', evidence_ids=['ev_1']))
    data['uwc_points']=points
    return data


class ContinuousCurveTests(unittest.TestCase):
    def test_a_b_dense_trace_reduced_and_same_x_aggregated(self):
        data=dense_trace(); process_continuous_traces(data)
        stats=data['curves'][0]['trace_processing']
        self.assertEqual(stats['raw_point_count'],402)
        self.assertEqual(stats['centerline_point_count'],201)
        self.assertLess(stats['exported_point_count'],60)
        self.assertEqual(len(data['raw_trace_points']),402)

    def test_c_d_transition_retained_and_flat_tail_not_oversampled(self):
        data=dense_trace(); process_continuous_traces(data)
        points=data['uwc_points']
        transition=sum(-3 <= p['temperature_celsius'] <= 0 for p in points)
        tail=sum(-20 <= p['temperature_celsius'] < -10 for p in points)
        self.assertGreaterEqual(transition,4)
        self.assertLess(tail,transition)

    def test_e_reconstruction_error_within_configured_tolerance(self):
        data=dense_trace(); process_continuous_traces(data)
        stats=data['curves'][0]['trace_processing']
        self.assertLessEqual(stats['raw_to_centerline_p95_normalized_error'],
                             stats['centerline_fidelity_p95_threshold_normalized'])
        self.assertLessEqual(stats['max_normalized_reconstruction_error'],
                             stats['simplification_tolerance_normalized']+1e-12)

    def test_foreign_parallel_trace_is_isolated_before_fidelity(self):
        data=dense_trace()
        extra=[]
        for point in data['uwc_points'][::2]:
            clone=copy.deepcopy(point)
            clone['point_id']='contaminated_'+clone['point_id']
            clone['uwc_original_value']+=.05
            clone['uwc_preferred_value']+=.05
            extra.append(clone)
        data['uwc_points'].extend(extra)
        for i,point in enumerate(data['uwc_points'],1):
            point['point_index']=point['trace_order']=i
        data['qa_settings']={'trace_centerline_fidelity_p95_threshold_normalized':.01}
        process_continuous_traces(data)
        stats=data['curves'][0]['trace_processing']
        self.assertLessEqual(stats['raw_to_centerline_p95_normalized_error'],.01)
        self.assertEqual(stats['excluded_contaminated_point_count'],201)
        self.assertIn('multimodal_path_candidate_warning',stats['warnings'])
        self.assertEqual(stats['curve_isolation_status'],'partial')
        self.assertEqual(stats['status'],'review_required')
        result=run(data)
        self.assertFalse(result['paper']['database_ready'])
        self.assertEqual(result['paper']['review_status'],'usable_with_review')

    def test_f_monotonicity_is_computed_on_scientific_curve(self):
        data=dense_trace()
        process_continuous_traces(data)
        stats=data['curves'][0]['trace_processing']
        self.assertEqual(stats['monotonicity_checked_on'],'scientific_curve_points')
        self.assertEqual(stats['monotonicity_violation_count'],0)

    def test_j_discrete_markers_are_untouched(self):
        data=graphical(); before=copy.deepcopy(data['uwc_points'])
        process_continuous_traces(data)
        self.assertEqual(data['uwc_points'],before)
        self.assertEqual(data.get('raw_trace_points'),[])

    def test_k_raw_trace_survives_simplified_excel_export(self):
        data=dense_trace()
        with tempfile.TemporaryDirectory() as tmp:
            result=export(data,tmp)
            saved=json.loads((Path(tmp)/'.internal'/'paper_test_UWC.json').read_text(encoding='utf8'))
            self.assertEqual(len(saved['raw_trace_points']),402)
            self.assertLess(len(saved['uwc_points']),60)
            book=load_workbook(Path(tmp)/'paper_test_UWC.xlsx',data_only=True)
            numeric=sum(isinstance(cell.value,(int,float)) for row in book['UWC_SFCC'] for cell in row)
            self.assertLess(numeric,150)

    def test_volumetric_trace_full_path_restores_preferred_theta_and_excel(self):
        data=dense_trace()
        with tempfile.TemporaryDirectory() as tmp:
            result=export(data,tmp)
            self.assertEqual(len(result['raw_trace_points']),402)
            self.assertLess(len(result['uwc_points']),60)
            self.assertTrue(all(p['uwc_preferred_basis']=='volumetric' and
                                p['uwc_preferred_unit']=='cubic_meter_per_cubic_meter' and
                                abs(p['uwc_preferred_value']-p['uwc_original_value'])<1e-12
                                for p in result['uwc_points']))
            book=load_workbook(Path(tmp)/'paper_test_UWC.xlsx',data_only=True)
            values=[cell.value for row in book['UWC_SFCC'] for cell in row]
            self.assertIn('θu (m³/m³)',values)
            self.assertTrue(any(abs(v-result['uwc_points'][0]['uwc_preferred_value'])<1e-12
                                for v in values if isinstance(v,(int,float))))

    def test_su_trace_full_path_converts_after_simplification(self):
        data=dense_trace();curve=data['curves'][0]
        curve.update(uwc_basis='unfrozen_water_saturation', raw_variable='unfrozen_water_saturation',
                     conversion_property_ids=['porosity_trace'])
        for point in data['uwc_points']:
            point.update(uwc_original_basis='unfrozen_water_saturation',uwc_original_unit='fraction',
                         uwc_original_unit_reported='-',uwc_preferred_value=None,uwc_preferred_unit=None,
                         uwc_preferred_basis=None,uwc_preferred_value_source=None,uwc_derivation_id=None)
        data['soil_properties']=[blank('property',property_id='porosity_trace',sample_id='s_1',condition_id='cond_1',
            property='porosity',property_reported='Porosity',property_role='porosity',experiment_role='uwc_nmr',
            applicability_evidence_ids=['ev_1'],value=.4,value_reported='0.40',unit='fraction',unit_reported='-',
            original_value=.4,original_unit='fraction',definition_reported='Pore volume / specimen volume',
            definition_status='explicit',value_source='reported',derivation_id=None,mineral_name_reported=None,
            mineral_name_normalized=None,confidence='high',evidence_ids=['ev_1'])]
        with tempfile.TemporaryDirectory() as tmp:
            result=export(data,tmp)
            self.assertTrue(result['derivations'])
            for point in result['uwc_points']:
                self.assertAlmostEqual(point['uwc_preferred_value'],.4*point['uwc_original_value'])
                self.assertEqual(point['uwc_preferred_basis'],'volumetric')
                record=next(d for d in result['derivations'] if d['derivation_id']==point['uwc_derivation_id'])
                self.assertTrue(record['scope_verified'])
                self.assertEqual(record['specimen_role'],'uwc_nmr')
            values=[cell.value for row in load_workbook(Path(tmp)/'paper_test_UWC.xlsx',data_only=True)['UWC_SFCC'] for cell in row]
            self.assertIn('θu (m³/m³)',values)


class PropertyRoleAndPrecisionTests(unittest.TestCase):
    def test_g_maximum_and_test_density_roles_are_distinct(self):
        data=graphical(); prop=data['soil_properties'][0]
        prop.update(property='maximum_dry_density', property_reported='Maximum dry density',
                    property_role='maximum_dry_density', experiment_role='compaction_test')
        self.assertEqual(issues(data),[])
        prop['property_role']='test_dry_density'
        self.assertTrue(any('incompatible property role' in e for e in issues(data)))

    def test_h_maximum_density_cannot_convert_gravimetric_uwc(self):
        data=dense_trace(); curve=data['curves'][0]
        curve.update(uwc_basis='gravimetric_dry_mass', raw_variable='gravimetric_unfrozen_water_content',
                     conversion_property_ids=['rho_1'])
        for point in data['uwc_points']:
            point.update(uwc_original_basis='gravimetric_dry_mass', uwc_original_unit='kilogram_per_kilogram',
                         uwc_original_unit_reported='kg/kg', uwc_preferred_value=None,
                         uwc_preferred_unit=None, uwc_preferred_basis=None, uwc_preferred_value_source=None)
        data['soil_properties']=[blank('property',property_id='rho_1',sample_id='s_1',condition_id=None,
            property='maximum_dry_density',property_reported='Maximum dry density',property_role='maximum_dry_density',
            experiment_role='compaction_test',applicability_evidence_ids=['ev_1'],value=1600,value_reported='1.60',
            unit='kilogram_per_cubic_meter',unit_reported='g/cm³',original_value=1.6,
            original_unit='gram_per_cubic_centimeter',definition_reported='Maximum dry density',definition_status='explicit',
            value_source='reported',derivation_id=None,mineral_name_reported=None,mineral_name_normalized=None,
            normalized=True,confidence='high',evidence_ids=['ev_1'])]
        result=run(data)
        self.assertEqual(len(result['raw_trace_points']),402)
        self.assertTrue(all(p['uwc_preferred_value'] is None for p in result['uwc_points']))
        self.assertFalse(result['paper']['database_ready'])

    def test_i_reported_condition_precision_survives_display(self):
        data=graphical()
        data['experimental_cases']=[dict(case_id='case_1',sample_id='s_1',case_label_reported='Precision case',
            varying_property='initial_volumetric_water_content',factor_label_reported='Initial water content',
            variables=[dict(property='initial_volumetric_water_content',values=[.1562,.2],unit='cubic_meter_per_cubic_meter',
                            values_reported=['0.1562','0.2000'],reported_precisions=[4,4],evidence_ids=['ev_1'])],
            condition_ids=['cond_1'],branches=['freezing'],curve_data_types=['measured_point'],
            temperature_range_celsius=[-20,0],evidence_ids=['ev_1'])]
        headers,rows=test_condition_table(data)
        self.assertIn('0.1562, 0.2000',rows[0])


@unittest.skipUnless(
    (ROOT/'fixtures'/'final_hardening').exists(),
    'Private recorded-paper fixtures are not distributed in the public repository.',
)
class RecordedPaperFixtureTests(unittest.TestCase):
    def read(self,name):
        return json.loads((ROOT/'fixtures'/'final_hardening'/name).read_text(encoding='utf8'))

    def test_sun_recorded_identity_and_dense_trace_contract(self):
        row=self.read('sun2023.json')
        self.assertEqual(row['identified_curve_count'],8)
        self.assertEqual(row['recorded_raw_trace_coordinate_count'],3012)
        self.assertEqual(row['representation'],'continuous_measured_trace')
        self.assertEqual(row['manual_reference_status'],'not_available')
        self.assertEqual(row['actual_processing_status'],'blocked_missing_structured_raw_trace')
        self.assertFalse(row['database_ready'])

    def test_cai_density_is_maximum_without_uwc_specimen_binding(self):
        row=self.read('cai2024.json')
        self.assertEqual(row['density']['property_role'],'maximum_dry_density')
        self.assertFalse(row['density']['uwc_specimen_binding_verified'])
        self.assertEqual(row['density']['reported_text'],'1.76')
        self.assertEqual(row['uwc_basis']['final_variable_classification'],'unknown_water_variable')
        self.assertFalse(row['uwc_basis']['conversion_performed'])
        self.assertFalse(row['database_ready'])

    def test_zou_discrete_no_regression_contract(self):
        row=self.read('zou2023.json')
        self.assertEqual((row['expected_curve_count'],row['recoverable_point_count']),(7,99))
        self.assertEqual(row['representation'],'discrete_markers')
        self.assertEqual(row['review_status'],'partial')
        self.assertFalse(row['interpolation_allowed'])
