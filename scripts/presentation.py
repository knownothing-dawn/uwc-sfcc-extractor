"""A per-paper human view; never mutate the normalized scientific records."""
from collections import Counter, defaultdict
import math
import re

from science import VOCAB, convert_value
from sources import selected_points
from semantics import data_semantics, representation


from coverage import curve_coverage
from design import (BASIC, STATE, basic_soil_table, test_condition_table, design_cases,
                    factor_value, factor_caption, case_label, soil_name, property_cell)

SOIL_COLUMNS, EXTRA_COLUMNS = BASIC, []

BRANCH_ORDER = {v: i for i, v in enumerate(("freezing", "thawing", "cooling", "warming", "unknown"))}
METHOD_LABELS = {"nuclear_magnetic_resonance": "NMR", "time_domain_reflectometry": "TDR",
                 "differential_scanning_calorimetry": "DSC"}
REVIEW_LABELS = {"axis_uncertain": "axis uncertain", "curve_identity_uncertain": "curve identity uncertain",
                 "legend_uncertain": "curve identity uncertain", "overlapping_curves": "digitization uncertain",
                 "digitization_uncertain": "digitization uncertain", "branch_uncertain": "branch uncertain",
                 "condition_uncertain": "condition uncertain", "anchor_conflict": "numerical source discrepancy",
                 "missing_point_or_not_reported": "point coverage incomplete",
                 "trace_processing_review": "continuous trace processing review"}


def unit_symbol(unit, reported=None):
    if unit in ("unknown", "other") and reported:
        return reported
    return VOCAB["units"].get(unit, {}).get("display_symbol", unit or "")


def sample_label(sample):
    return sample["sample_name_reported"] or sample["soil_type_reported"] or sample["soil_type_normalized"]


def condition_title(condition, properties):
    return condition["condition_label_reported"] or "Experimental condition"


soil_table = basic_soil_table


def point_pairs(points, branch):
    """Use recorded sequence, then a deterministic thermal ordering as fallback."""
    indices = [p.get("point_index") for p in points]
    sequenced = all(isinstance(i, int) and not isinstance(i, bool) for i in indices) and len(set(indices)) == len(indices)
    ordered = sorted(points, key=lambda p: p["point_index"]) if sequenced else sorted(
        points, key=lambda p: p["temperature_celsius"], reverse=branch in ("freezing", "cooling"))
    preferred = all(p["uwc_preferred_value"] is not None and p["uwc_preferred_basis"] == "volumetric"
                    and p["uwc_preferred_unit"] == "cubic_meter_per_cubic_meter" for p in ordered)
    if preferred:
        return "θu (m³/m³)", [(p["temperature_celsius"], p["uwc_preferred_value"]) for p in ordered]
    # A block must have one honest axis unit. Never mix raw mass ratios with θ.
    first = ordered[0]
    if first["uwc_original_basis"] in ("unfrozen_water_saturation", "liquid_water_saturation"):
        basis = first["uwc_original_basis"]
        if any(p["uwc_original_basis"] != basis for p in ordered):
            raise ValueError("Incompatible saturation definitions within one curve")
        label = "Su" if basis == "unfrozen_water_saturation" else "Sl"
        try:
            return f"{label} (-)", [(p["temperature_celsius"], convert_value(
                p["uwc_original_value"], p["uwc_original_unit"], "fraction")) for p in ordered]
        except ValueError:
            if all((p["uwc_original_unit"],p["uwc_original_unit_reported"]) ==
                   (first["uwc_original_unit"],first["uwc_original_unit_reported"]) for p in ordered):
                return f"{label} ({unit_symbol(first['uwc_original_unit'],first['uwc_original_unit_reported'])}; review unit)", [
                    (p["temperature_celsius"],p["uwc_original_value"]) for p in ordered]
            raise ValueError("Incompatible saturation units within one curve")
    unit, basis = first["uwc_original_unit"], first["uwc_original_basis"]
    pairs = []
    for point in ordered:
        source = point["uwc_original_unit"]
        if point["uwc_original_basis"] != basis:
            raise ValueError("A curve has incompatible UWC bases; resolve the record before export")
        if source == unit and (unit not in ("other", "unknown") or
                              point["uwc_original_unit_reported"] == first["uwc_original_unit_reported"]):
            value = point["uwc_original_value"]
        else:
            value = convert_value(point["uwc_original_value"], source, unit, basis)
        pairs.append((point["temperature_celsius"], value))
    label = "θ" if basis == "volumetric" else "w_u, dry mass" if basis == "gravimetric_dry_mass" else "UWC, " + basis.replace("_", " ")
    return f"{label} ({unit_symbol(unit, first['uwc_original_unit_reported'])})", pairs


def curve_blocks(package):
    """The block plan is internal; visible cells never contain technical IDs."""
    samples = {s["sample_id"]: i for i, s in enumerate(package["samples"])}
    conditions = {c["condition_id"]: c for c in package["conditions"]}
    condition_order = {cid: i for i, cid in enumerate(conditions)}
    points = defaultdict(list)
    for point in selected_points(package):
        points[point["curve_id"]].append(point)
    reviews = defaultdict(set)
    for review in package["review_queue"]:
        if review["status"] == "open" and review["issue_type"] in REVIEW_LABELS:
            reviews[review["object_id"]].add(REVIEW_LABELS[review["issue_type"]])
    blocks = []
    coverage = package.get("coverage") or curve_coverage(package)
    visible = set(coverage["unique_curve_ids"])
    curves = sorted([c for c in package["curves"] if c["curve_id"] in visible or not points[c["curve_id"]]], key=lambda c: (samples[c["sample_id"]], condition_order[c["condition_id"]],
                                                     BRANCH_ORDER.get(c["branch"], 99)))
    duplicate_labels = Counter((c["sample_id"], c["condition_id"], c["branch"]) for c in curves)
    repeated = Counter()
    anchored = {p['curve_id'] for p in package.get('reference_points', []) if p['point_id'] in
                {s['reference_point_id'] for s in package.get('observation_selections', [])}}
    case_plan = design_cases(package)
    cycles_by_sample = defaultdict(set)
    for condition in conditions.values():
        if condition["cycle_number"] is not None:
            cycles_by_sample[condition["sample_id"]].add(condition["cycle_number"])
    for curve in curves:
        cid, sid, state = curve["curve_id"], curve["sample_id"], curve["condition_id"]
        observations = points[cid]
        heading, pairs = point_pairs(observations, curve["branch"]) if observations else ("UWC (unit unresolved)", [])
        related = {package["paper"]["paper_id"], cid, sid, state, *curve["evidence_ids"]}
        for point in observations:
            related.update([point["point_id"], *point["evidence_ids"]])
        flags = set().union(*(reviews[x] for x in related))
        if curve["branch"] in ("unknown", "other"):
            flags.add("branch uncertain")
        if conditions[state]["confidence"] == "low":
            flags.add("condition uncertain")
        digitization = curve["digitization"]
        if digitization and digitization["axis_confidence"] != "high":
            flags.add("axis uncertain")
        if curve["confidence"] == "low" and curve["extraction_method"] not in ("table_direct", "text_direct"):
            flags.add("digitization uncertain")
        if any(p["confidence"] == "low" and p["value_source"] == "figure_digitized" for p in observations):
            flags.add("digitization uncertain")
        subtitle = curve["branch"].replace("_", " ").capitalize()
        cycles = cycles_by_sample[sid]
        if len(cycles) > 1 and curve["cycle_number"] is not None and not any(
                c["varying_property"] == "cycle_number" and state in c["condition_ids"] for c in case_plan):
            subtitle += f" · Cycle {curve['cycle_number']}"
        if curve.get('data_semantics') and curve['curve_data_type']=='unknown':
            subtitle += " · " + data_semantics(curve).capitalize() + " · " + representation(curve).replace('_', ' ')
        elif curve["curve_data_type"] != "measured_point":
            subtitle += " · " + curve["curve_data_type"].replace("_", " ")
        if curve["data_origin"] != "original_data":
            subtitle += " · " + curve["data_origin"].replace("_", " ")
        if duplicate_labels[(sid, state, curve["branch"])] > 1:
            subtitle += " · " + (curve["legend_label"] or METHOD_LABELS.get(curve["method"], curve["method"].replace("_", " ")))
        title = condition_title(conditions[state], package["soil_properties"])
        label = (sid, title, subtitle)
        repeated[label] += 1
        if repeated[label] > 1:
            subtitle += f" · Series {repeated[label]}"
        notes = (["Review required: " + "; ".join(sorted(flags)) + "."] if flags else [])
        if cid in anchored: notes.append("Exact numerical anchors included; see report.")
        blocks.append({"curve_id": cid, "sample_id": sid, "condition_id": state, "title": title, "subtitle": subtitle,
                       "note": " ".join(notes) or None,
                       "uwc_heading": heading, "pairs": pairs})
    return blocks


def workbook(package, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "UWC_SFCC"
    basic_headers, basic_rows = basic_soil_table(package)
    test_headers, test_rows = test_condition_table(package)
    columns = max(8, len(basic_headers), len(test_headers))
    widths = [18] + [14] * (columns - 1)
    for headers, rows in ((basic_headers, basic_rows), (test_headers, test_rows)):
        for col in range(len(headers)):
            content = max((len(str(row[col] or "")) for row in rows), default=0)
            widths[col] = max(widths[col], min(20 if headers[col] == "Soil" else 18, content + 2))
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    palette = {"paper": "29445B", "soil": "EEDFAA", "test": "EEDFAA", "section": "DFEAF5",
               "group": "E8EFDE", "curve": "F2F2F2"}
    for name in ("paper", "soil", "test", "section", "group", "curve", "text", "number", "temperature", "warning", "label"):
        is_header = name in palette
        font = Font(name="Calibri", size=14 if name == "paper" else 11 if name in ("section", "group", "soil", "test") else 10,
                    bold=is_header or name == "label", color="FFFFFF" if name == "paper" else "9C3E10" if name == "warning" else "222222")
        style = NamedStyle(name=name, font=font, alignment=Alignment(
            horizontal="left" if name in ("text", "warning", "label") else "center", vertical="center", wrap_text=name not in ("number", "temperature")),
            number_format="0.00" if name == "temperature" else "0.0000" if name == "number" else "General")
        if is_header:
            style.fill = PatternFill("solid", fgColor=palette[name])
        wb.add_named_style(style)
    thin = Side(style="thin", color="D8D8D8")
    medium = Side(style="medium", color="BAC8D2")
    pointer = f".internal/{package['paper']['paper_id']}_UWC.json"

    def put(row, col, value, style=None, fmt=None):
        if isinstance(value, str) and (len(value.encode("utf-16-le")) // 2 > 32000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", value)):
            value = f"[Original content preserved in {pointer}]"
        cell = ws.cell(row, col, value)
        if isinstance(value, str):cell.data_type = "s"
        cell.style = style or ("number" if isinstance(value, (float, int)) else "text")
        if fmt:cell.number_format = fmt
        return cell

    def span(row, start, end, value, style="text", merge=True):
        for col in range(start, end + 1):
            put(row, col, value if col == start else None, style)
        if merge and end > start:
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        size = sum(widths[start - 1:end]) - 3
        count = sum(2 if ord(c) > 255 else 1 for c in str(value or ""))
        lines = max(str(value or "").count("\n") + 1, math.ceil(count / max(size, 1)))
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 18,
                                            23 if style in ("paper", "soil", "test", "section", "group") else 17,
                                            lines * 14 + 4)

    def section(row, title, style):
        span(row, 1, columns, title, style)
        for col in range(1, columns + 1):
            ws.cell(row, col).border = Border(bottom=medium)

    def table(row, headers, rows, style):
        header_lines = max(math.ceil(len(h) / (widths[i] - 2)) for i, h in enumerate(headers))
        ws.row_dimensions[row].height = max(30, 14 * header_lines + 6)
        for col, header in enumerate(headers, 1):
            cell = put(row, col, header, "curve")
            cell.border = Border(bottom=thin)
        for i, values in enumerate(rows, row + 1):
            height = 18
            for col, value in enumerate(values, 1):
                fmt = "General" if isinstance(value,str) else "0.00" if "(%)" in headers[col - 1] else "0.000"
                put(i, col, value, fmt=fmt)
                if isinstance(value, str):
                    height = max(height, 14 * math.ceil(sum(2 if ord(c) > 255 else 1 for c in value) / (widths[col - 1] - 2)) + 4)
            ws.row_dimensions[i].height = height
        return row + max(1, len(rows)) + 1

    paper = package["paper"]
    section(1, "Paper Information", "paper")
    put(2, 1, "Paper Title", "label")
    span(2, 2, columns, paper["paper_title"])
    put(3, 1, "Authors", "label")
    span(3, 2, columns, "; ".join(paper["authors"]))
    put(4, 1, "Year", "label");put(4, 2, paper["publication_year"], fmt="0")
    put(4, 3, "Journal", "label");span(4, 4, 5, paper["journal"])
    put(4, 6, "DOI", "label");span(4, 7, columns, paper["doi"])
    methods = list(dict.fromkeys(METHOD_LABELS.get(c["method"], c["method_reported"] or c["method"].replace("_", " ")) for c in package["curves"]))
    put(5, 1, "Method", "label");put(5, 2, "; ".join(methods))
    put(5, 3, "SFCC Branch", "label");span(5, 4, 5, "; ".join(dict.fromkeys(c["branch"].capitalize() for c in package["curves"])))
    sources = []
    for evidence in package["evidence"]:
        if evidence.get("source_role") == "primary_uwc":
            sources.append("Fig. " + evidence["figure_number"] if evidence["figure_number"] else
                           "Table " + evidence["table_number"] if evidence["table_number"] else
                           evidence["source_file"].replace("\\", "/").split("/")[-1])
    put(5, 6, "Primary source", "label");span(5, 7, columns, "; ".join(dict.fromkeys(sources)))
    coverage = package.get("coverage") or curve_coverage(package)
    percentage = f"{coverage['coverage_percentage']:.2f}%" if coverage["coverage_percentage"] is not None else "not inventoried"
    from quality import review_severity
    critical = sum(r['status']=='open' and review_severity(r)=='critical' for r in package['review_queue'])
    readiness = "DATABASE READY" if paper.get('database_ready') else "NOT DATABASE READY"
    axis = package.get('numerical_quality', {}).get('axis_metrics', [])
    axis_text = (f"axis ticks {sum(m['x_anchor_count'] + m['y_anchor_count'] for m in axis)} "
                 f"({sum(m['status']=='pass' for m in axis)}/{len(axis)} mappings verified)" if axis else "axis ticks not checked")
    gates = package.get('qa_gates', {})
    vnext = ", ".join(f"{name.replace('_',' ')}={gates.get(name,'not_checked')}"
                      for name in ('detection_completeness','numerical_quality','physical_quality') if name in gates)
    put(6, 1, "Coverage / QA", "label")
    span(6, 2, columns, f"Curves {coverage['matched_curve_count']} / {coverage['expected_curve_count']} ({percentage}); "
         f"points {package.get('point_coverage', {}).get('extracted_point_count', len(package['uwc_points']))}; "
         f"{axis_text}; cross-source anchors {package.get('anchor_summary', {}).get('checked_count', 0)}; "
         f"{vnext}; critical issues {critical}. {readiness}. See report.",
         "text" if paper.get('database_ready') else "warning")
    from copy import copy
    status_font = copy(ws.cell(6,2).font)
    status_font.color = '9C0006' if critical else '27632A' if paper.get('database_ready') else '9C3E10'
    ws.cell(6,2).font = status_font
    section(8, "Basic Soil Properties", "soil")
    row = table(9, basic_headers, basic_rows, "soil") + 1
    section(row, "Test Conditions", "test")
    row = table(row + 1, test_headers, test_rows, "test") + 1
    uwc_start = row
    section(row, "UWC/SFCC Data", "section");row += 1
    blocks = curve_blocks(package)
    by_condition = defaultdict(list)
    for block in blocks:by_condition[block["condition_id"]].append(block)
    samples = {s["sample_id"]: soil_name(s) for s in package["samples"]}
    cases = design_cases(package)
    factors = list(dict.fromkeys(c["varying_property"] for c in cases))
    used = set()

    def write_band(band, start):
        note = any(b["note"] for b in band)
        header = start + 2 + int(note)
        for i, block in enumerate(band):
            col = 2 * i + 1
            span(start, col, col + 1, block["title"], "curve")
            span(start + 1, col, col + 1, block["subtitle"], "curve")
            if note and block["note"]:span(start + 2, col, col + 1, block["note"], "warning")
            for j, label in enumerate(("Temp (°C)", block["uwc_heading"])):
                cell = put(header, col + j, label, "curve");cell.border = Border(bottom=thin)
                ws.row_dimensions[header].height = max(ws.row_dimensions[header].height or 20, 14 * math.ceil(len(label)/(widths[col+j-1]-2)) + 4)
            for index, (temperature, value) in enumerate(block["pairs"], header + 1):
                put(index, col, temperature, "temperature");put(index, col + 1, value)
                ws.row_dimensions[index].height = 17
            if not block["pairs"]:span(header + 1, col, col + 1, "Missing data; see report", "warning")
        return header + max(1, max(len(b["pairs"]) for b in band)) + 3

    for factor in factors:
        factor_cases = [c for c in cases if c["varying_property"] == factor]
        if factor != "none":
            section(row, "Effect of " + case_label(factor_cases[0]), "section");row += 1
        for case in factor_cases:
            other_factors = set(factors) - {factor, "none"}
            fixed = [factor_caption(v["property"], v["values"][0], v["unit"],
                                    (v.get("reported_precisions") or [None])[0],
                                    (v.get("values_reported") or [None])[0]) for v in case["variables"]
                     if v["property"] in other_factors and len(v["values"]) == 1]
            section(row, samples[case["sample_id"]] + (" — " + "; ".join(fixed) if fixed else ""), "group");row += 1
            selected, reused = [], []
            for cid in case["condition_ids"]:
                for original in by_condition[cid]:
                    block = dict(original)
                    value = factor_value(package, cid, factor)
                    if value is not None:block["title"] = factor_caption(factor, *value)
                    if block["curve_id"] in used:
                        if block["title"] not in reused:reused.append(block["title"])
                        continue
                    selected.append(block);used.add(block["curve_id"])
            missing = [m for m in coverage["missing_conditions"] if m["condition_id"] in case["condition_ids"]]
            if reused:
                span(row, 1, min(8, columns), "Shown above: " + "; ".join(reused));row += 1
            if missing:
                labels=[]
                for m in missing:
                    value=factor_value(package,m["condition_id"],factor)
                    labels.append(factor_caption(factor,*value) if value else "Unrecovered curve")
                span(row, 1, min(8, columns), "Missing: " + "; ".join(dict.fromkeys(labels)) + ". See report.", "warning");row += 1
            for offset in range(0, len(selected), 4):
                row = write_band(selected[offset:offset + 4], row)
    # Unexpected curves remain visible and flagged; no hidden loss during grouping.
    remaining = [b for b in blocks if b["curve_id"] not in used]
    if remaining:
        section(row, "Additional extracted curves — review condition binding", "group");row += 1
        for offset in range(0, len(remaining), 4):row = write_band(remaining[offset:offset + 4], row)
    if not blocks:span(row, 1, columns, "No UWC points recovered; see report.", "warning")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    if ws.max_row > 80:
        ws.freeze_panes = "A7"  # Retain only the small paper block, leaving summaries visible while scrolling.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = ws.dimensions
    wb.save(path)
