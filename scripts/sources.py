"""Bounded numerical-source recovery and exact anchors; no paper-specific rules."""
import copy
import csv
import hashlib
import json
from collections import defaultdict
from pathlib import Path

from manifest import branch_key, settings, dataset_identity
from science import convert_value, normalize_point
from semantics import data_semantics, representation

PRIORITY = {'structured_data': 0, 'table': 1, 'numerical_label': 2, 'numerical_text': 3}
METHODS = {'structured_data': 'supplementary_direct', 'table': 'table_direct',
           'numerical_label': 'text_direct', 'numerical_text': 'text_direct'}
VARIABLES = {'volumetric': 'volumetric_unfrozen_water_content',
    'gravimetric_dry_mass': 'gravimetric_unfrozen_water_content',
    'unfrozen_water_saturation': 'unfrozen_water_saturation',
    'liquid_water_saturation': 'liquid_water_saturation', 'normalized': 'normalized_water_content',
    'mass_fraction_wet_mass': 'mass_fraction'}


def resolve_variables(package):
    for curve in package['curves']:
        # This records an already evidence-assigned basis, never infers from magnitude.
        curve.setdefault('raw_variable', VARIABLES.get(curve['uwc_basis'], 'unknown_water_variable'))
        curve.setdefault('raw_variable_reported', curve['uwc_definition_reported'])
        curve.setdefault('variable_definition_verified', False)


def load_datasets(package, artifact_root):
    """Read each explicitly bound local table once. Never search unrelated files."""
    cache, digests, sheet_rows, results, failures = {}, {}, {}, [], []
    bounds = {}
    for source in package.get('numerical_datasets', []):
        b = source.get('table_binding')
        if not b: continue
        path = Path(b['path'])
        if not path.is_absolute(): path = Path(artifact_root) / path
        key = (str(path.resolve()), b['sheet'])
        old = bounds.get(key, (b['row_start'], b['row_end']))
        bounds[key] = (min(old[0],b['row_start']), max(old[1],b['row_end']))
    for source in package.get('numerical_datasets', []):
        dataset = copy.deepcopy(source)
        binding = dataset.get('table_binding')
        try:
            if binding:
                path = Path(binding['path'])
                if not path.is_absolute(): path = Path(artifact_root) / path
                key = str(path.resolve())
                if key not in digests: digests[key] = hashlib.sha256(path.read_bytes()).hexdigest()
                if digests[key] != binding['sha256']:
                    raise ValueError('Bound numerical file changed; source verification is required')
                if key not in cache:
                    if path.suffix.lower() == '.xlsx':
                        from openpyxl import load_workbook
                        cache[key] = load_workbook(path, read_only=True, data_only=True)
                    elif path.suffix.lower() in ('.csv', '.tsv'):
                        with path.open(encoding='utf-8-sig', newline='') as stream:
                            cache[key] = list(csv.reader(stream, delimiter='\t' if path.suffix.lower() == '.tsv' else ','))
                    else: raise ValueError('Only explicitly bound XLSX/CSV/TSV numerical tables are supported')
                table = cache[key]
                if binding['row_start'] > binding['row_end']:
                    raise ValueError('Invalid numerical table row bounds')
                if hasattr(table, 'sheetnames'):
                    sheet_key = (key,binding['sheet'])
                    start,end = bounds[sheet_key]
                    if sheet_key not in sheet_rows:
                        sheet_rows[sheet_key] = list(table[binding['sheet']].iter_rows(min_row=start,max_row=end,values_only=True))
                    rows = sheet_rows[sheet_key][binding['row_start']-start:binding['row_end']-start+1]
                else: rows = table[binding['row_start']-1:binding['row_end']]
                observations = []
                for row in rows:
                    t, value = row[binding['temperature_column']-1], row[binding['value_column']-1]
                    if t is None and value is None: continue
                    if isinstance(t, bool) or isinstance(value, bool): raise ValueError('Boolean in numerical observation')
                    observations.append(dict(temperature=float(t), temperature_unit=binding['temperature_unit'],
                        value=float(value), unit=binding['unit'], temperature_reported=str(t), value_reported=str(value)))
                dataset['observations'] = observations
            results.append(dataset)
        except (OSError, ValueError, TypeError, KeyError, IndexError) as exc:
            failures.append(dict(dataset_id=source['dataset_id'], message=str(exc)))
    for table in cache.values():
        if hasattr(table, 'close'): table.close()
    package['source_failures'] = failures
    return results


def source_verified(dataset, evidence):
    return (dataset['variable_definition_verified'] and dataset['evidence_ids'] and
            VARIABLES.get(dataset['uwc_basis']) == dataset['raw_variable'] and
            all(eid in evidence and evidence[eid]['transcription_status'] == 'verified' for eid in dataset['evidence_ids']))


def new_point(dataset, observation, curve_id, point_id, index):
    from quality import SCHEMA
    point = {k: None for k in SCHEMA['$defs']['point']['required']}
    method = METHODS[dataset['source_kind']]
    point.update(point_id=point_id, curve_id=curve_id, point_index=index,
        temperature_original_value=observation['temperature'], temperature_original_unit=observation['temperature_unit'],
        temperature_original_unit_reported=observation['temperature_unit'], temperature_value_source='reported',
        uwc_original_value=observation['value'], uwc_original_unit=observation['unit'],
        uwc_original_unit_reported=observation['unit'], uwc_original_basis=dataset['uwc_basis'],
        curve_data_type=dataset['curve_data_type'], extraction_method=method, value_source='reported',
        evidence_ids=dataset['evidence_ids'], confidence='high')
    return normalize_point(point)


def targeted_search(package, expected, artifact_root):
    evidence = {e['evidence_id']: e for e in package['evidence']}
    selected = [evidence[eid] for eid in expected['expected_source_ids'] if eid in evidence]
    pages = sorted({e['page'] for e in selected if e.get('page')})
    condition = next(c for c in package['conditions'] if c['condition_id'] == expected['condition_id'])
    terms = [condition['condition_label_reported'] or ''] + [f"{x['value']:g}" for x in expected['condition_values']]
    terms = [t for t in dict.fromkeys(terms) if t]
    texts = {}
    cache = Path(artifact_root) / '.internal/candidates.json'
    if cache.exists():
        content = json.loads(cache.read_text(encoding='utf-8'))
        if content.get('source_sha256') == package['paper']['source_sha256']:
            texts = {p['page']: p['text_reported'] for p in content['pages'] if p['page'] in pages}
    if not texts:
        pdfs = {e['source_file'] for e in selected if Path(e['source_file']).suffix.lower() == '.pdf' and Path(e['source_file']).is_file()}
        if pdfs:
            import pymupdf
            for path in sorted(pdfs):
                with pymupdf.open(path) as doc:
                    for number in pages:
                        if 1 <= number <= len(doc): texts[number] = doc[number-1].get_text()
    hits = []
    for page, text in texts.items():
        for line in text.splitlines():
            if any(term.casefold() in line.casefold() for term in terms):
                hits.append(f'p. {page}: ' + line)
    return pages, terms, hits[:12]


def recover_missing(package, datasets, artifact_root):
    """One targeted retry per unchanged source/expectation, never a whole-PDF rerun."""
    from coverage import curve_coverage, condition_keys, observation_fingerprint
    from quality import SCHEMA
    keys = condition_keys(package)
    evidence = {e['evidence_id']: e for e in package['evidence']}
    coverage = curve_coverage(package)
    missing = {dataset_identity(m, keys) for m in coverage['missing_conditions']}
    attempts = {a['expected_curve_id']: a for a in package.get('recovery_attempts', [])}
    source_digest = hashlib.sha256(json.dumps([datasets, package.get('source_failures', [])], sort_keys=True).encode()).hexdigest()
    for expected in package.get('expected_datasets', []):
        key = dataset_identity(expected, keys)
        if key not in missing: continue
        fingerprint = hashlib.sha256(json.dumps([expected, source_digest], sort_keys=True).encode()).hexdigest()
        old = attempts.get(expected['expected_curve_id'])
        if old and old['input_fingerprint'] == fingerprint and old['status'] != 'recovered': continue
        candidates = [d for d in datasets if dataset_identity(d, keys) == key and d['coverage_role'] == 'full_series'
            and d['observations'] and source_verified(d, evidence)
            and any(evidence[e]['source_role'] == 'primary_uwc' for e in d['evidence_ids'])]
        pages, terms, hits = [], [], []
        status, detail = 'not_found', 'No verified full-series numerical source matches the missing scientific condition.'
        if candidates:
            priority = min(PRIORITY[d['source_kind']] for d in candidates)
            candidates = [d for d in candidates if PRIORITY[d['source_kind']] == priority]
            fingerprints = {(d['method'], d['data_origin'], observation_fingerprint([new_point(d, p, 'temporary', 'temporary', i+1)
                for i, p in enumerate(d['observations'])])) for d in candidates}
            if len(fingerprints) > 1:
                status, detail = 'ambiguous', 'Equal-priority full sources disagree; no source was silently selected.'
            else:
                source = candidates[0]
                cid = 'recovered_' + expected['expected_curve_id']
                if any(c['curve_id'] == cid for c in package['curves']):
                    status, detail = 'ambiguous', 'Recovery record already exists without usable observations; inspect its binding.'
                else:
                    condition = next(c for c in package['conditions'] if c['condition_id'] == expected['condition_id'])
                    curve = {k: None for k in SCHEMA['$defs']['curve']['required']}
                    curve.update(curve_id=cid, paper_id=package['paper']['paper_id'], sample_id=expected['sample_id'],
                        condition_id=condition['condition_id'], cycle_number=condition['cycle_number'], branch=source['branch'],
                        branch_reported=source['branch'], branch_source='reported', method=source['method'], method_reported=source['method'],
                        curve_data_type=source['curve_data_type'], uwc_basis=source['uwc_basis'],
                        raw_variable=source['raw_variable'], raw_variable_reported=source.get('variable_definition_reported'),
                        variable_definition_verified=True, uwc_definition_reported=source.get('variable_definition_reported'),
                        temperature_unit='degree_celsius', extraction_method=METHODS[source['source_kind']],
                        data_origin=source['data_origin'], original_citation_reported=source.get('original_citation_reported'),
                        needs_manual_review=False, confidence='high', evidence_ids=source['evidence_ids'],
                        source_points_verified=True, legend_label=condition['condition_label_reported'],
                        data_semantics=data_semantics(source), representation=representation(source),
                        source_occurrences=[dict(evidence_ids=source['evidence_ids'], case_id=None,
                            legend_label=condition['condition_label_reported'], relationship='primary', source_locator=source['source_locator'])])
                    package['curves'].append(curve)
                    if source.get('series_variant_reported'): curve['series_variant_reported'] = source['series_variant_reported']
                    package['uwc_points'].extend(new_point(source, p, cid, cid + '_p' + str(i+1), i+1) for i, p in enumerate(source['observations']))
                    status, detail = 'recovered', 'Recovered the verified full numerical series from its explicitly bound source.'
        else:
            pages, terms, hits = targeted_search(package, expected, artifact_root)
            if pages: status = 'targeted_review_required'
            if package.get('source_failures'): status = 'source_unavailable'
        attempts[expected['expected_curve_id']] = dict(expected_curve_id=expected['expected_curve_id'],
            input_fingerprint=fingerprint, status=status, source_ids=expected['expected_source_ids'],
            target_pages=pages, search_terms=terms, source_hits=hits, detail=detail)
    package['recovery_attempts'] = list(attempts.values())


def selected_points(package):
    """Already resolved upstream: exporters perform no scientific source selection."""
    references = {p['point_id']: p for p in package.get('reference_points', [])}
    selections = {s['point_id']: references[s['reference_point_id']] for s in package.get('observation_selections', [])}
    return [dict(selections[p['point_id']], point_id=p['point_id'], point_index=p['point_index'])
            if p['point_id'] in selections else p for p in package['uwc_points']]


def comparison_fingerprint(check, point, reference):
    """Invalidate an earlier source decision when either observation changes."""
    return hashlib.sha256(json.dumps([check, point, reference], sort_keys=True).encode()).hexdigest()


def validate_anchors(package, datasets):
    from coverage import condition_keys
    from conversion import convert_preferred
    keys = condition_keys(package)
    evidence = {e['evidence_id']: e for e in package['evidence']}
    limits = settings(package)
    grouped, points = defaultdict(list), defaultdict(list)
    for c in package['curves']: grouped[dataset_identity(c, keys)].append(c)
    for point in package['uwc_points']: points[point['curve_id']].append(point)
    checks, references, choices = [], [], defaultdict(list)
    old_reference_ids = {p['point_id'] for p in package.get('reference_points', [])}
    package['derivations'] = [d for d in package['derivations'] if not any(i['source_id'] in old_reference_ids for i in d['input_variables'])]
    for source in datasets:
        if not source_verified(source, evidence):
            package.setdefault('source_failures', []).append(dict(dataset_id=source['dataset_id'],
                message='Numerical source variable/evidence has not been verified: '+source['source_locator']))
            continue
        for index, observation in enumerate(source['observations'], 1):
            check = dict(check_id='anchor_' + source['dataset_id'] + '_' + str(index), dataset_id=source['dataset_id'],
                observation_index=index, curve_id=None, point_id=None, reference_point_id=None,
                absolute_error=None, normalized_error=None, comparison_unit=None, temperature_error_celsius=None,
                status='unmatched', selected=False, detail='No unique observation matches material, condition, branch and temperature.')
            matching = [c for c in grouped[dataset_identity(source, keys)] if c['method'] == source['method'] and c['data_origin'] == source['data_origin']]
            # A direct primary series does not need to validate itself.
            matching = [c for c in matching if c['extraction_method'] in ('pdf_vector', 'raster_marker', 'raster_curve')]
            if not matching:
                if source['coverage_role'] == 'anchors_only': checks.append(check)
                continue
            if len(matching) != 1:
                check.update(status='ambiguous', detail='Multiple curves match the anchor identity; replicate/condition resolution is required.')
                checks.append(check); continue
            curve = matching[0];check['curve_id'] = curve['curve_id']
            if source['uwc_basis'] != curve['uwc_basis'] or source['raw_variable'] != curve.get('raw_variable'):
                check.update(status='incompatible_variable', detail='Anchor and curve have different physical variables; no magnitude-based matching or conversion.')
                checks.append(check); continue
            t = convert_value(observation['temperature'], observation['temperature_unit'], 'degree_celsius')
            cp = [p for p in points[curve['curve_id']] if abs(p['temperature_celsius'] - t) <= limits['temperature_match_tolerance_celsius']]
            if len(cp) != 1:
                check['status'] = 'ambiguous' if cp else 'unmatched';checks.append(check);continue
            point = cp[0]
            reference = new_point(source, observation, curve['curve_id'], 'reference_' + source['dataset_id'] + '_' + str(index), index)
            temporary = dict(package, curves=[curve], uwc_points=[reference], derivations=[])
            converted = convert_preferred(temporary, copy_package=False)
            reference = converted['uwc_points'][0]
            package['derivations'].extend(converted['derivations'])
            references.append(reference)
            unit = 'cubic_meter_per_cubic_meter' if point['uwc_preferred_value'] is not None and reference['uwc_preferred_value'] is not None else point['uwc_original_unit']
            if unit == 'cubic_meter_per_cubic_meter':
                a, b = point['uwc_preferred_value'], reference['uwc_preferred_value']
            else:
                a = point['uwc_original_value'];b = convert_value(reference['uwc_original_value'], reference['uwc_original_unit'], unit, curve['uwc_basis'])
            error = abs(a-b)
            span = curve.get('plotted_uwc_axis_span')
            # Use a declared, unit-bearing axis span; never estimate it from sampled values.
            try: span = convert_value(span, curve['plotted_uwc_axis_unit'], unit, curve['uwc_basis']) if span else None
            except ValueError: span = None
            normalized = error/span if span and span > 0 else None
            strong = unit == 'cubic_meter_per_cubic_meter' and error > limits['volumetric_strong_threshold'] + 1e-12
            warning = (unit == 'cubic_meter_per_cubic_meter' and error > limits['volumetric_warning_threshold'] + 1e-12 or
                       normalized is not None and normalized > limits['normalized_warning_threshold'] + 1e-12)
            check.update(point_id=point['point_id'], reference_point_id=reference['point_id'], absolute_error=error,
                normalized_error=normalized, comparison_unit=unit, temperature_error_celsius=abs(point['temperature_celsius']-t),
                status='strong_warning' if strong else 'warning' if warning else 'within_tolerance',
                detail='Original graphical observation retained; higher-priority exact observation is eligible for the output pair.')
            choices[point['point_id']].append((PRIORITY[source['source_kind']], reference, check))
            checks.append(check)
    selections = []
    for pid, candidates in choices.items():
        priority = min(x[0] for x in candidates)
        candidates = [x for x in candidates if x[0] == priority]
        values = {(p['temperature_celsius'], p['uwc_preferred_value'] if p['uwc_preferred_value'] is not None else
                   convert_value(p['uwc_original_value'], p['uwc_original_unit'], candidates[0][1]['uwc_original_unit'], p['uwc_original_basis']))
                  for _, p, _ in candidates}
        if len(values) > 1:
            for _, _, check in candidates:
                check.update(status='conflicting_references', detail='Equal-priority explicit values disagree; no averaging or automatic selection.')
            continue
        _, reference, check = candidates[0]
        check['selected'] = True
        selections.append(dict(point_id=pid, reference_point_id=reference['point_id']))
    package.update(reference_points=references, observation_selections=selections, anchor_checks=checks)
